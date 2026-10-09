from __future__ import annotations

from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import cast
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook

from larp_bot.adapters.memory import MemoryEventRepository, MemoryRegistrationRepository, MemoryUserRepository
from larp_bot.adapters.yandex_disk.repository import PUBLIC_HEADERS, YandexDiskShowcaseRepository
from larp_bot.adapters.ydb.repositories import YdbExecutor, YdbUserRepository
from larp_bot.application.services import RegistrationCatalog
from larp_bot.domain.models import AttendanceStatus, Event, Platform, Registration, TelegramUser, VkUser
from larp_bot.domain.security import participant_key
from tests.conftest import MemoryDiskStore

SECRET = "pass-column-test-secret"
EVENT_ID = "b49df060-85f7-45fd-851a-2e3dcf706320"


class CountingUsers(MemoryUserRepository):
    def __init__(self) -> None:
        super().__init__()
        self.preference_reads = 0

    async def list_pass_preferences(self):
        self.preference_reads += 1
        return await super().list_pass_preferences()

    async def list_all(self):
        raise AssertionError("table refresh must not load complete profiles")

    async def get(self, platform, user_id):
        raise AssertionError("table refresh must not issue a query for each player")

    async def find_linked(self, user):
        raise AssertionError("links must be resolved from the same batch")


@pytest.mark.asyncio
@pytest.mark.parametrize("deferred", [False, True])
async def test_admin_pass_answers_refresh_main_and_reserve_without_changing_public_table(disk_store, deferred):
    now = datetime.now(UTC)
    users = CountingUsers()
    profiles = [
        TelegramUser(tg_id=1, needs_pass=True),
        VkUser(vk_id=2, needs_pass=False),
        TelegramUser(tg_id=3, needs_pass=None),
    ]
    for user in profiles:
        await users.save(user)
    rows = [
        Registration(
            event_id=EVENT_ID,
            participant_key=participant_key(SECRET, platform, user_id, EVENT_ID),
            display_name="Игрок Тестовый",
            wish_play="Без пожеланий",
            created_at=now + timedelta(seconds=index),
        )
        for index, (platform, user_id) in enumerate(
            ((Platform.TELEGRAM, 1), (Platform.VK, 2), (Platform.TELEGRAM, 3), (Platform.VK, 99))
        )
    ]
    event = Event(
        event_id=EVENT_ID,
        name="Тест пропусков",
        game_date=now.date(),
        player_amount=1,
        disk_resource_path="disk:/larp-bot/events/master.xlsx",
        public_registration_url="https://disk.example/master",
        public_table_resource_path="disk:/larp-bot/events/public.xlsx",
        public_table_public_url="https://disk.example/public",
        registrations_migrated_at=now,
    )
    events = MemoryEventRepository()
    await events.create(event, [])
    uploader = None
    if deferred:
        uploader = AsyncMock()

        async def submit(files, **kwargs):
            await disk_store.submit(files, **kwargs)
            return None

        uploader.submit.side_effect = submit
    catalog = RegistrationCatalog(
        events,
        MemoryRegistrationRepository(rows),
        YandexDiskShowcaseRepository(disk_store, uploader),
        users=users,
        participant_secret=SECRET,
    )
    if event.disk_resource_path not in disk_store.files:
        await catalog.showcase.create_event_workbook(event.disk_resource_path, player_amount=event.player_amount)
        if event.public_table_resource_path is not None:
            await catalog.showcase.create_public_event_workbook(
                event.public_table_resource_path, player_amount=event.player_amount
            )
    await catalog.refresh(event)
    assert users.preference_reads == 1
    workbook = load_workbook(BytesIO(disk_store.files[event.disk_resource_path]))
    try:
        sheet = workbook.active
        assert sheet["J1"].value == "Нужен ли пропуск"
        assert [sheet.cell(row, 10).value for row in (2, 4, 5, 6)] == ["Да", "Нет", "Не указано", "Не указано"]
        assert sheet.auto_filter.ref == "A1:J1"
        assert str(sheet.merged_cells) == "A3:J3"
    finally:
        workbook.close()
    public_before = load_workbook(BytesIO(disk_store.files[event.public_table_resource_path]))
    public_values = list(public_before.active.values)
    assert public_values[0] == PUBLIC_HEADERS
    public_before.close()
    profiles[0].needs_pass = False
    await users.save(profiles[0])
    await catalog.refresh(event)
    assert users.preference_reads == 2
    workbook = load_workbook(BytesIO(disk_store.files[event.disk_resource_path]))
    assert workbook.active["J2"].value == "Нет"
    workbook.close()
    public_after = load_workbook(BytesIO(disk_store.files[event.public_table_resource_path]))
    assert list(public_after.active.values) == public_values
    public_after.close()
    assert await catalog.registrations.list_for_event(EVENT_ID) == rows


@pytest.mark.asyncio
@pytest.mark.parametrize("reciprocal,ambiguous", [(True, False), (False, False), (True, True)])
async def test_pass_column_uses_latest_reciprocal_global_profile(disk_store, reciprocal, ambiguous):
    now = datetime.now(UTC)
    users = CountingUsers()
    await users.save(
        TelegramUser(
            tg_id=1,
            telegram_handle="@player",
            linked_vk_id=2,
            needs_pass=False,
            updated_at=now,
        )
    )
    await users.save(
        VkUser(
            vk_id=2,
            telegram_handle="@player" if reciprocal else "@someone_else",
            needs_pass=True,
            updated_at=now + timedelta(seconds=1),
        )
    )
    if ambiguous:
        await users.save(TelegramUser(tg_id=3, telegram_handle="@player", linked_vk_id=2))
    event = Event(
        event_id=EVENT_ID,
        name="Тест",
        game_date=now.date(),
        player_amount=1,
        disk_resource_path="disk:/larp-bot/events/master.xlsx",
        public_registration_url="https://disk.example/master",
        public_table_resource_path="disk:/larp-bot/events/public.xlsx",
        public_table_public_url="https://disk.example/public",
        registrations_migrated_at=now,
    )
    rows = [
        Registration(
            event_id=EVENT_ID,
            participant_key=participant_key(SECRET, platform, user_id, EVENT_ID),
            display_name="Игрок Тестовый",
            wish_play="Без пожеланий",
            created_at=now + timedelta(seconds=index),
        )
        for index, (platform, user_id) in enumerate(((Platform.TELEGRAM, 1), (Platform.VK, 2)))
    ]
    catalog = RegistrationCatalog(
        MemoryEventRepository(),
        MemoryRegistrationRepository(rows),
        YandexDiskShowcaseRepository(disk_store),
        users=users,
        participant_secret=SECRET,
    )
    if event.disk_resource_path not in disk_store.files:
        await catalog.showcase.create_event_workbook(event.disk_resource_path, player_amount=event.player_amount)
        if event.public_table_resource_path is not None:
            await catalog.showcase.create_public_event_workbook(
                event.public_table_resource_path, player_amount=event.player_amount
            )
    await catalog.refresh(event)
    workbook = load_workbook(BytesIO(disk_store.files[event.disk_resource_path]))
    assert workbook.active["J2"].value == ("Да" if reciprocal and not ambiguous else "Нет")
    assert workbook.active["J4"].value == "Да"
    workbook.close()
    assert users.preference_reads == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("cancelled", [False, True])
async def test_empty_or_cancelled_roster_does_not_read_profiles(disk_store: MemoryDiskStore, event: Event, cancelled):
    users = CountingUsers()
    event = event.model_copy(update={"public_table_resource_path": "disk:/larp-bot/events/public.xlsx"})
    rows = (
        [
            Registration(
                event_id=event.event_id,
                participant_key="a" * 43,
                display_name="Игрок",
                wish_play="Любой",
                attendance_status=AttendanceStatus.CANCELLED,
            )
        ]
        if cancelled
        else []
    )
    catalog = RegistrationCatalog(
        MemoryEventRepository(),
        MemoryRegistrationRepository(rows),
        YandexDiskShowcaseRepository(disk_store),
        users=users,
        participant_secret=SECRET,
    )
    if event.disk_resource_path not in disk_store.files:
        await catalog.showcase.create_event_workbook(event.disk_resource_path, player_amount=event.player_amount)
        if event.public_table_resource_path is not None:
            await catalog.showcase.create_public_event_workbook(
                event.public_table_resource_path, player_amount=event.player_amount
            )
    await catalog.refresh(event)
    assert users.preference_reads == 0


@pytest.mark.asyncio
async def test_ydb_pass_projection_reads_only_small_fields_in_two_queries():
    now = datetime.now(UTC)
    db = AsyncMock()

    async def query(sql, params, *, read_only):
        assert params == {"$after": 0}
        assert read_only
        assert "SELECT" in sql
        assert not any(field in sql for field in ("pass_details_json", "dialog_context_json", "full_name", "email"))
        if "tg_users" in sql:
            return [
                {"tg_id": 1, "needs_pass": False, "telegram_handle": "@player", "linked_vk_id": 2, "updated_at": now}
            ]
        return [{"vk_id": 2, "needs_pass": True, "telegram_handle": "@player", "updated_at": now}]

    db.query.side_effect = query
    profiles = await YdbUserRepository(cast(YdbExecutor, db)).list_pass_preferences()
    assert db.query.call_count == 2
    assert [(row.platform, row.user_id, row.needs_pass) for row in profiles] == [
        (Platform.TELEGRAM, 1, False),
        (Platform.VK, 2, True),
    ]


@pytest.mark.asyncio
async def test_editing_linked_vk_profile_refreshes_telegram_registration(disk_store, event):
    from larp_bot.adapters.memory import MemoryCommandPublisher
    from larp_bot.application.services import OrderedMutationService, RegistrationService

    now = datetime.now(UTC)
    event = event.model_copy(
        update={
            "public_table_resource_path": "disk:/larp-bot/events/public.xlsx",
            "registrations_migrated_at": now,
        }
    )
    telegram = TelegramUser(
        tg_id=1,
        full_name="Игрок Тестовый",
        vk_url="https://vk.com/id2",
        crossplay=False,
        telegram_handle="@player",
        linked_vk_id=2,
        needs_pass=False,
        updated_at=now,
    )
    vk = VkUser(vk_id=2, telegram_handle="@player", needs_pass=True, updated_at=now + timedelta(seconds=1))
    users = MemoryUserRepository()
    await users.save(telegram)
    await users.save(vk)
    row = Registration(
        event_id=event.event_id,
        participant_key=participant_key(SECRET, Platform.TELEGRAM, 1, event.event_id),
        display_name="Игрок Тестовый",
        wish_play="Без пожеланий",
    )
    events = MemoryEventRepository([event])
    showcase = YandexDiskShowcaseRepository(disk_store)
    await showcase.create_event_workbook(event.disk_resource_path)
    await showcase.create_public_event_workbook(event.public_table_resource_path)
    catalog = RegistrationCatalog(
        events,
        MemoryRegistrationRepository([row]),
        showcase,
        users=users,
        participant_secret=SECRET,
    )
    await catalog.refresh(event)
    before = load_workbook(BytesIO(disk_store.files[event.disk_resource_path]))
    assert before.active["J2"].value == "Да"
    before.close()
    vk.needs_pass = False
    vk.updated_at = now + timedelta(seconds=2)
    await users.save(vk)
    publisher = MemoryCommandPublisher()
    service = RegistrationService(events, catalog, publisher, SECRET, users)
    assert await service.enqueue_profile_updates(Platform.VK, 2, idempotency_key="edit-pass") == 1
    command = publisher.commands[0]
    assert command.platform is Platform.TELEGRAM
    assert command.participant_key == row.participant_key
    await OrderedMutationService(events, catalog, users, SECRET).apply(command)
    after = load_workbook(BytesIO(disk_store.files[event.disk_resource_path]))
    assert after.active["J2"].value == "Нет"
    after.close()


@pytest.mark.asyncio
async def test_ydb_pass_projection_pages_through_all_users():
    now = datetime.now(UTC)
    db = AsyncMock()
    records = [{"tg_id": user_id, "needs_pass": True, "updated_at": now} for user_id in range(1, 1202)]

    async def query(sql, params, *, read_only):
        assert read_only
        assert "ORDER BY" in sql and "LIMIT 500" in sql
        if "vk_users" in sql:
            return []
        return [row for row in records if row["tg_id"] > params["$after"]][:500]

    db.query.side_effect = query
    profiles = await YdbUserRepository(cast(YdbExecutor, db)).list_pass_preferences()
    assert [profile.user_id for profile in profiles] == list(range(1, 1202))
    assert db.query.call_count == 4
