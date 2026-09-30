from __future__ import annotations

from io import BytesIO

import pypdfium2 as pdfium
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN

from larp_bot.adapters.yandex_disk.achievements import (
    ACHIEVEMENT_FONT,
    AchievementTemplateError,
    _fitted_font_size,
    achievement_presentation_bytes,
)
from tests.conftest import achievement_template_bytes


def blank_pdf_bytes(page_count: int) -> bytes:
    document = pdfium.PdfDocument.new()
    try:
        for _ in range(page_count):
            document.new_page(100, 100).close()
        output = BytesIO()
        document.save(output)
        return output.getvalue()
    finally:
        document.close()


def test_achievement_renderer_converts_pdf_and_replaces_every_centered_placeholder() -> None:
    template = achievement_template_bytes()
    original = bytes(template)
    game_name = "Очень длинное\nназвание игры для проверки"

    rendered = achievement_presentation_bytes(template, game_name)

    assert template == original
    presentation = Presentation(BytesIO(rendered))
    assert len(presentation.slides) == 1
    assert presentation.slide_width.pt == pytest.approx(841)
    assert presentation.slide_height.pt == pytest.approx(595)
    slide = presentation.slides[0]
    pictures = [shape for shape in slide.shapes if shape.shape_type is MSO_SHAPE_TYPE.PICTURE]
    text_boxes = [shape for shape in slide.shapes if shape.has_text_frame]
    assert len(pictures) == 1
    assert pictures[0].left == pictures[0].top == 0
    assert pictures[0].width == presentation.slide_width
    assert pictures[0].height == presentation.slide_height
    assert len(text_boxes) == 8
    background = Image.open(BytesIO(pictures[0].image.blob)).convert("RGB")
    preserved_blue_details = (
        (255, 737),
        (892, 737),
        (1503, 737),
        (2123, 737),
        (258, 1608),
        (879, 1609),
        (1494, 1608),
        (2123, 1609),
    )
    assert all(background.getpixel(point) == (47, 84, 150) for point in preserved_blue_details)

    normalized_name = "Очень длинное название игры для проверки"
    assert {shape.text for shape in text_boxes} == {normalized_name}
    expected_columns = (112, 321, 526, 734)
    expected_rows = (260, 550)
    for shape in text_boxes:
        center_x = shape.left.pt + shape.width.pt / 2
        center_y = shape.top.pt + shape.height.pt / 2
        assert any(center_x == pytest.approx(column, abs=5) for column in expected_columns)
        assert any(center_y == pytest.approx(row, abs=5) for row in expected_rows)
    for shape in text_boxes:
        frame = shape.text_frame
        paragraph = frame.paragraphs[0]
        run = paragraph.runs[0]
        assert shape.fill.type is MSO_FILL_TYPE.BACKGROUND
        assert shape.line.fill.type is MSO_FILL_TYPE.BACKGROUND
        assert frame.auto_size is MSO_AUTO_SIZE.NONE
        assert frame.word_wrap is False
        assert frame.vertical_anchor is MSO_ANCHOR.MIDDLE
        assert paragraph.alignment is PP_ALIGN.CENTER
        assert run.font.name == ACHIEVEMENT_FONT == "Oswald"
        assert run.font.size is not None
        assert run.font.size.pt < 15


def test_achievement_renderer_keeps_source_font_size_when_the_name_fits() -> None:
    rendered = achievement_presentation_bytes(achievement_template_bytes(), "Игра")

    presentation = Presentation(BytesIO(rendered))
    sizes = {
        shape.text_frame.paragraphs[0].runs[0].font.size.pt
        for shape in presentation.slides[0].shapes
        if shape.has_text_frame
    }
    assert sizes == {15.0}


@pytest.mark.parametrize(
    ("game_name", "expected_size"),
    [
        ("И" * 20, 15.0),
        ("И" * 21, 14.25),
        ("И" * 40, 7.5),
        ("И" * 300, 1.0),
    ],
)
def test_achievement_font_size_scales_by_game_name_character_count(
    game_name: str,
    expected_size: float,
) -> None:
    assert _fitted_font_size(game_name, 15.0) == expected_size


def test_achievement_renderer_requires_the_single_template_page() -> None:
    with pytest.raises(AchievementTemplateError, match="exactly one page"):
        achievement_presentation_bytes(blank_pdf_bytes(2), "Игра")


def test_achievement_renderer_requires_a_placeholder() -> None:
    with pytest.raises(AchievementTemplateError, match="ЛОГОС"):
        achievement_presentation_bytes(blank_pdf_bytes(1), "Игра")


def test_achievement_renderer_rejects_an_invalid_pdf() -> None:
    with pytest.raises(AchievementTemplateError, match="valid PDF"):
        achievement_presentation_bytes(b"not a pdf", "Игра")
