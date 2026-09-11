from __future__ import annotations

from collections.abc import Sequence
from datetime import UTC, date, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from larp_bot.adapters.memory import (
    MemoryCommandPublisher,
    MemoryDeferredTransport,
    MemoryEventRepository,
    MemoryRegistrationRepository,
    MemoryUserRepository,
    MemoryVkUserIdResolver,
    StaticAdminProvider,
)
from larp_bot.adapters.yandex_disk.repository import YandexDiskShowcaseRepository
from larp_bot.application.conversation import (
    ADD_EVENT_LEADER,
    ADD_SELF_EVENT_LEADER,
    ADMIN,
    ADMIN_MENU,
    ARCHIVE_GAME,
    ARCHIVED_GAMES,
    BACK,
    CANCEL,
    CHANGE_PLAYER_AMOUNT,
    CHANGE_STATUS,
    CHARACTER_WISHES,
    CO_PLAYER_WISH_SKIP,
    CO_PLAYER_WISHES,
    CONFIRM,
    CREATION_MASTER_QUESTION,
    ENLIST,
    EVENT_TABLES,
    FREE_TEXT_DIALOG_STATES,
    GAME_WISHES,
    GAMEMASTER_NOTIFICATION,
    GAMEMASTER_REVOKED_NOTIFICATION,
    GRANT_GAMEMASTER,
    INACTIVE_GAME_RESPONSE,
    INVALID_INPUT_RESPONSE,
    LIST_EVENT_LEADERS,
    LIST_GAMEMASTERS,
    LIST_WAITING_PLAYERS,
    MAIN_MENU,
    MAIN_MENU_LABEL,
    MANAGE_GAMES,
    MANAGE_MASTERS,
    MANAGE_MODERATORS,
    MANAGE_NOTIFICATIONS,
    MANAGE_PLAYERS,
    MASTER_TABLE_DISCLAIMER,
    MY_GAMES,
    PAGE_BACK,
    PAGE_FORWARD,
    PAGE_SIZE,
    PROFILE,
    REMOVE_EVENT_LEADER,
    REMOVE_PLAYER,
    REMOVE_WAITING_PLAYERS,
    REVOKE_GAMEMASTER,
    SEND_CONFIRMATION_REMINDER,
    SEND_CONFIRMED_LINK,
    SEND_CONFIRMED_NOTIFICATION,
    STATUS_CLOSED,
    STATUS_CONFIRMATION,
    STATUS_REGISTRATION,
    UNARCHIVE_GAME,
    UNSUPPORTED_CONTENT_RESPONSE,
    VENUE_ADDRESS,
    ConversationEngine,
)
from larp_bot.application.services import (
    EventAdministrationService,
    RegistrationCatalog,
    RegistrationService,
)
from larp_bot.domain.models import (
    AttendanceStatus,
    BotIdentity,
    Button,
    CharacterWishPayload,
    ConfirmationDeadlinePayload,
    CoPlayerWishPayload,
    EnlistPayload,
    Event,
    EventStatus,
    InboundMessage,
    NotificationPayload,
    Operation,
    PassDetails,
    Platform,
    PlayerAmountPayload,
    ProfileSyncPayload,
    Registration,
    TelegramUser,
    VkUser,
)
from tests.conftest import MemoryDiskStore


def inbound(
    update: int,
    value: str,
    *,
    callback: bool = False,
    user_id: int = 1,
    telegram_username: str | None = None,
) -> InboundMessage:
    return InboundMessage(
        identity=BotIdentity(platform=Platform.TELEGRAM, platform_user_id=user_id),
        update_id=str(update),
        text="" if callback else value,
        callback=value if callback else None,
        chat_id=user_id,
        telegram_username=telegram_username,
    )


class FailOnceTransport(MemoryDeferredTransport):
    def __init__(self) -> None:
        super().__init__()
        self.failed = False

    async def send(
        self,
        *,
        platform: Platform,
        user_id: int,
        request_id: str,
        text: str,
        buttons: Sequence[Button] = (),
    ) -> None:
        if not self.failed:
            self.failed = True
            raise RuntimeError("temporary transport failure")
        await super().send(
            platform=platform,
            user_id=user_id,
            request_id=request_id,
            text=text,
            buttons=buttons,
        )


async def engine_setup(
    store: MemoryDiskStore,
    event: Event,
    *additional_events: Event,
    admin: bool = False,
    gamemaster: bool = False,
    tg_admin_ids: set[int] | None = None,
    vk_admin_ids: set[int] | None = None,
    vk_profile_ids: dict[str, int] | None = None,
) -> tuple[
    ConversationEngine,
    MemoryUserRepository,
    MemoryCommandPublisher,
    MemoryRegistrationRepository,
]:
    users = MemoryUserRepository()
    await users.save(
        TelegramUser(
            tg_id=1,
            full_name="Иван Иванов",
            vk_url="https://vk.com/id1",
            crossplay=True,
            larp_experience=True,
            needs_pass=False,
        )
    )
    events = MemoryEventRepository([event, *additional_events])
    tables = MemoryRegistrationRepository()
    showcase = YandexDiskShowcaseRepository(store)
    await showcase.create_event_workbook(event.disk_resource_path)
    catalog = RegistrationCatalog(events, tables, showcase)
    publisher = MemoryCommandPublisher()
    registrations = RegistrationService(events, catalog, publisher, "participant-secret")
    transport = MemoryDeferredTransport()
    conversation = ConversationEngine(
        users,
        events,
        registrations,
        EventAdministrationService(events, showcase, catalog, users, "participant-secret"),
        StaticAdminProvider(
            tg_ids=({1} if admin else set()) | (tg_admin_ids or set()),
            vk_ids=vk_admin_ids,
            tg_gamemaster_ids={1} if gamemaster else set(),
        ),
        transport=transport,
        vk_user_ids=MemoryVkUserIdResolver(vk_profile_ids),
    )
    return conversation, users, publisher, tables


@pytest.mark.asyncio
async def test_unsupported_content_is_rejected_without_advancing_dialog(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event)
    await engine.handle(inbound(1, PROFILE))
    await engine.handle(inbound(2, "profile:edit", callback=True))
    before = await users.get(Platform.TELEGRAM, 1)
    assert before is not None
    assert before.dialog_state == "PROFILE_FULL_NAME_CYRILLIC"

    response = await engine.handle(
        InboundMessage(
            identity=BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1),
            update_id="3",
            has_unsupported_content=True,
            chat_id=1,
        )
    )
    after = await users.get(Platform.TELEGRAM, 1)

    assert response.text == UNSUPPORTED_CONTENT_RESPONSE
    assert after is not None
    assert after.dialog_state == before.dialog_state
    assert after.dialog_context == before.dialog_context
    assert after.last_bot_buttons == before.last_bot_buttons


@pytest.mark.asyncio
async def test_enlist_finishes_without_confirmation_or_character_wishes(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, _, publisher, _ = await engine_setup(disk_store, event)
    responses = [await engine.handle(inbound(1, ENLIST))]
    responses.append(await engine.handle(inbound(2, f"select:enlist:{event.event_id}", callback=True)))
    responses.append(await engine.handle(inbound(3, "Алиса", telegram_username="ivan_player")))
    prompts = " ".join(response.text.casefold() for response in responses)
    assert "пожелания по персонажу" not in prompts
    assert "подтвердить запись" not in prompts
    assert not any(button.value == "enlist:confirm" for response in responses for button in response.buttons)
    assert len(publisher.commands) == 1
    queued = publisher.commands[0]
    assert queued.operation is Operation.ENLIST
    assert isinstance(queued.payload, EnlistPayload)
    assert not hasattr(queued.payload, "character_wish")
    assert queued.payload.larp_experience is True
    assert queued.payload.crossplay is True
    assert queued.payload.vk_profile == "https://vk.com/id1"
    assert queued.payload.telegram_profile == "https://t.me/ivan_player"
    assert [(button.label, button.value) for button in queued.reply_context.buttons] == [(MAIN_MENU_LABEL, MAIN_MENU)]
    assert [(button.label, button.value) for button in responses[-1].buttons] == [(MAIN_MENU_LABEL, MAIN_MENU)]
    assert responses[-1].deferred


@pytest.mark.asyncio
async def test_cancel_confirmation_warns_about_losing_queue_position(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, _, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Алиса",
    )

    games = await engine.handle(inbound(1, CANCEL))
    assert any(button.value == f"select:cancel:{event.event_id}" for button in games.buttons)
    response = await engine.handle(inbound(2, f"select:cancel:{event.event_id}", callback=True))

    assert response.text == (
        f"Вы уверены, что хотите отменить участие в игре «{event.name}»? "
        "Если вы отмените участие, то сможете записаться снова, но попадёте "
        "в конец очереди на основной состав."
    )
    assert [button.label for button in response.buttons] == ["Да, отменить", "Нет, оставить участие"]


@pytest.mark.asyncio
async def test_only_idle_buttonless_user_responses_get_main_menu_navigation(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, _, _, _ = await engine_setup(disk_store, event)

    prompt = await engine.handle(inbound(1, PROFILE, user_id=2))
    terminal = await engine.handle(inbound(2, "select:character:missing-game", callback=True))

    assert all(button.value != MAIN_MENU for button in prompt.buttons)
    assert [(button.label, button.value) for button in terminal.buttons] == [(MAIN_MENU_LABEL, MAIN_MENU)]


@pytest.mark.asyncio
async def test_terminal_navigation_buttons_leave_active_text_dialogs(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, admin=True)
    await engine.handle(inbound(1, PROFILE, user_id=2))

    main = await engine.handle(inbound(2, MAIN_MENU, callback=True, user_id=2))
    user = await users.get(Platform.TELEGRAM, 2)

    assert PROFILE in [button.value for button in main.buttons]
    assert user is not None and user.dialog_state == "IDLE"

    await engine.handle(inbound(3, "➕ Создать игру", callback=True))

    administration = await engine.handle(inbound(4, ADMIN_MENU, callback=True))
    admin = await users.get(Platform.TELEGRAM, 1)

    assert administration.text == "🛠 Администрирование"
    assert admin is not None and admin.dialog_state == "IDLE"


@pytest.mark.asyncio
async def test_enlist_rejects_stale_game_button_and_offers_explicit_skip(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, _, publisher, _ = await engine_setup(disk_store, event)
    await engine.handle(inbound(1, ENLIST))
    prompt = await engine.handle(inbound(2, f"select:enlist:{event.event_id}", callback=True))
    assert prompt.text == (
        "Есть ли игроки, с которыми вы хотели бы играть в одной сессии? "
        "Укажите каждого как «Фамилия Имя», через запятую. "
        "Мастера постараются учесть пожелание, но не могут его гарантировать.\n\n"
        "Ответ должен быть не длиннее 2000 символов. Если пожеланий нет, нажмите «Пропустить»."
    )
    assert "нажмите «пропустить»" in prompt.text.casefold()
    assert any(button.value == "enlist:wish-play:skip" for button in prompt.buttons)

    stale = await engine.handle(inbound(3, f"select:enlist:{event.event_id}", callback=True))
    assert "текущую клавиатуру" in stale.text
    assert "Если вы увидели это сообщение, попробуйте написать боту /start" in stale.text
    assert not publisher.commands

    response = await engine.handle(inbound(4, "enlist:wish-play:skip", callback=True))
    assert response.deferred
    assert isinstance(publisher.commands[-1].payload, EnlistPayload)
    assert publisher.commands[-1].payload.wish_play == "Без пожеланий"


@pytest.mark.asyncio
async def test_overlong_co_player_wish_is_rejected_without_leaving_the_prompt(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, users, publisher, _ = await engine_setup(disk_store, event)
    await engine.handle(inbound(1, ENLIST))
    prompt = await engine.handle(inbound(2, f"select:enlist:{event.event_id}", callback=True))

    rejected = await engine.handle(inbound(3, "А" * 2001))
    saved = await users.get(Platform.TELEGRAM, 1)

    assert rejected.text == INVALID_INPUT_RESPONSE
    assert rejected.buttons == prompt.buttons
    assert saved is not None and saved.dialog_state == "ENLIST_WISH_PLAY"
    assert publisher.commands == []


@pytest.mark.asyncio
async def test_single_registered_game_opens_wishes_menu_and_edits_co_player_wish(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, users, publisher, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Без пожеланий",
    )
    await tables.confirm(
        event.event_id,
        operation_id="existing-confirm",
        participant_key=key,
        character_wish="Doctor",
    )

    menu = await engine.handle(inbound(1, GAME_WISHES))
    prompt = await engine.handle(inbound(2, f"wishes:co-player:{event.event_id}", callback=True))
    accepted = await engine.handle(inbound(3, "Петров Пётр"))
    saved = await users.get(Platform.TELEGRAM, 1)

    assert menu.text == f"Пожелания по игре «{event.name}»"
    assert menu.buttons == [
        Button(label=CHARACTER_WISHES, value=f"wishes:character:{event.event_id}"),
        Button(label=CO_PLAYER_WISHES, value=f"wishes:co-player:{event.event_id}"),
    ]
    assert prompt.text == (
        "Есть ли игроки, с которыми вы хотели бы играть в одной сессии? "
        "Укажите каждого как «Фамилия Имя», через запятую. "
        "Мастера постараются учесть пожелание, но не могут его гарантировать.\n\n"
        "Ответ должен быть не длиннее 2000 символов. Если пожеланий нет, нажмите «Пропустить»."
    )
    assert prompt.buttons[0] == Button(label="Пропустить", value=CO_PLAYER_WISH_SKIP)
    assert accepted.deferred and accepted.command_enqueued
    assert saved is not None and saved.dialog_state == "IDLE"
    command = publisher.commands[-1]
    assert command.operation is Operation.UPDATE_CO_PLAYER_WISH
    assert isinstance(command.payload, CoPlayerWishPayload)
    assert command.payload.wish_play == "Петров Пётр"


@pytest.mark.asyncio
async def test_co_player_wish_edit_can_clear_wishes_with_skip(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, _, publisher, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Петров Пётр",
    )

    await engine.handle(inbound(1, GAME_WISHES))
    await engine.handle(inbound(2, f"wishes:co-player:{event.event_id}", callback=True))
    response = await engine.handle(inbound(3, CO_PLAYER_WISH_SKIP, callback=True))

    assert response.deferred
    command = publisher.commands[-1]
    assert command.operation is Operation.UPDATE_CO_PLAYER_WISH
    assert isinstance(command.payload, CoPlayerWishPayload)
    assert command.payload.wish_play == "Без пожеланий"


@pytest.mark.asyncio
@pytest.mark.parametrize("confirmed", [False, True])
async def test_active_registration_cannot_enlist_again(
    confirmed: bool,
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, _, publisher, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Без пожеланий",
    )
    if confirmed:
        await tables.confirm(
            event.event_id,
            operation_id="existing-confirm",
            participant_key=key,
            character_wish="Doctor",
        )

    await engine.handle(inbound(1, ENLIST))
    rejected = await engine.handle(inbound(2, f"select:enlist:{event.event_id}", callback=True))

    assert rejected.text == f"Вы уже записаны на игру «{event.name}»."
    assert publisher.commands == []


@pytest.mark.asyncio
async def test_duplicate_update_does_not_enqueue_twice(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, publisher, _ = await engine_setup(disk_store, event)
    await engine.handle(inbound(1, ENLIST))
    await engine.handle(inbound(2, f"select:enlist:{event.event_id}", callback=True))
    final = inbound(3, "Алиса")
    await engine.handle(final)
    duplicate = await engine.handle(final)
    assert duplicate.silent
    assert len(publisher.commands) == 1


@pytest.mark.asyncio
async def test_confirmation_is_first_character_wish_prompt(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, publisher, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Алиса",
    )
    listing = await engine.handle(inbound(1, CONFIRM))
    assert any(button.value == f"select:confirm:{event.event_id}" for button in listing.buttons)
    selection = await engine.handle(inbound(2, f"select:confirm:{event.event_id}", callback=True))
    assert "пожелания" in selection.text.casefold()
    stale = await engine.handle(inbound(3, f"select:confirm:{event.event_id}", callback=True))
    assert "кнопка устарела" in stale.text.casefold()
    assert stale.buttons == selection.buttons
    assert not publisher.commands
    await engine.handle(inbound(4, "Doctor"))
    command = publisher.commands[-1]
    assert command.operation is Operation.CONFIRM
    assert command.reply_context.text_success == ""


@pytest.mark.asyncio
async def test_signup_wish_prompt_warns_about_another_registration_on_the_same_day(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    game_date = date(2026, 11, 14)
    event.game_date = game_date
    other = event.model_copy(
        update={
            "event_id": "event-b1",
            "name": "Другая игра",
            "disk_resource_path": "disk:/larp-bot/events/event-b1-other.xlsx",
            "public_registration_url": "https://disk.example/public/other",
        }
    )
    engine, _, _, tables = await engine_setup(disk_store, event, other)
    other_key = engine.registrations.key(Platform.TELEGRAM, 1, other.event_id)
    await tables.enlist(
        other.event_id,
        operation_id="other-enlist",
        participant_key=other_key,
        display_name="Иван Иванов",
        wish_play="Без пожеланий",
    )

    prompt = await engine.handle(inbound(1, f"select:enlist:{event.event_id}", callback=True))

    assert prompt.text.startswith(
        "Важно: вы уже записаны на другую игру, которая проводится 14.11.26. "
        "Вы сможете подтвердить своё участие только в одной игре, когда решите, "
        "в какую именно хотите играть.\n\nЕсть ли игроки"
    )


@pytest.mark.asyncio
async def test_same_day_confirmation_requires_replacement_consent_before_character_wishes(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    event.game_date = date(2026, 11, 14)
    other = event.model_copy(
        update={
            "event_id": "event-b1",
            "name": "Другая игра",
            "disk_resource_path": "disk:/larp-bot/events/event-b1-other.xlsx",
            "public_registration_url": "https://disk.example/public/other",
        }
    )
    engine, users, publisher, tables = await engine_setup(disk_store, event, other)
    target_key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    other_key = engine.registrations.key(Platform.TELEGRAM, 1, other.event_id)
    for selected_event, key, operation_id in (
        (event, target_key, "target-enlist"),
        (other, other_key, "other-enlist"),
    ):
        await tables.enlist(
            selected_event.event_id,
            operation_id=operation_id,
            participant_key=key,
            display_name="Иван Иванов",
            wish_play="Без пожеланий",
        )
    await tables.confirm(
        other.event_id,
        operation_id="other-confirm",
        participant_key=other_key,
        character_wish="Врач",
    )

    replacement = await engine.handle(inbound(1, f"select:confirm:{event.event_id}", callback=True))

    assert replacement.text == (
        "Вы уже подтвердили своё участие в игре «Другая игра», которая проводится 14.11.26. "
        "Чтобы подтвердить своё участие в «Лесной предел» вам придётся отменить своё участие "
        "в «Другая игра». Вы уверены, что хотите это сделать?"
    )
    assert [button.label for button in replacement.buttons] == [
        "Да, я хочу участвовать в этой игре, а не в другой",
        "Нет, не хочу ничего менять",
    ]
    saved = await users.get(Platform.TELEGRAM, 1)
    assert saved is not None and saved.dialog_state == "CONFIRM_REPLACE"
    assert not publisher.commands

    wishes = await engine.handle(inbound(2, "confirm:replace:yes", callback=True))
    assert "пожелания по персонажу" in wishes.text.casefold()
    queued = await engine.handle(inbound(3, "Детектив"))

    assert queued.deferred
    command = publisher.commands[-1]
    assert command.operation is Operation.CONFIRM
    assert command.queue_group_id == "confirmation:telegram:1"
    assert command.payload == CharacterWishPayload(
        character_wish="Детектив",
        replaced_event_id=other.event_id,
    )


@pytest.mark.asyncio
async def test_overlong_character_wish_is_rejected_without_confirming(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, users, publisher, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Без пожеланий",
    )
    await engine.handle(inbound(1, CONFIRM))
    prompt = await engine.handle(inbound(2, f"select:confirm:{event.event_id}", callback=True))

    rejected = await engine.handle(inbound(3, "А" * 3001))
    saved = await users.get(Platform.TELEGRAM, 1)

    assert rejected.text == INVALID_INPUT_RESPONSE
    assert rejected.buttons == prompt.buttons
    assert saved is not None and saved.dialog_state == "CONFIRM_WISH"
    assert publisher.commands == []


@pytest.mark.asyncio
async def test_start_always_resets_dialog_and_returns_main_menu(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, users, publisher, _ = await engine_setup(disk_store, event)
    stale_button = Button(label="Устаревшая кнопка", value="stale:button")

    for update, state in enumerate(sorted(FREE_TEXT_DIALOG_STATES), start=1):
        user = await users.get(Platform.TELEGRAM, 1)
        assert user is not None
        user.dialog_state = state
        user.dialog_context = {"sentinel": state}
        user.last_bot_buttons = [stale_button]
        await users.save(user)

        message = inbound(update, "/start").model_copy(update={"has_unsupported_content": True})
        response = await engine.handle(message)
        saved = await users.get(Platform.TELEGRAM, 1)

        assert response.text == (
            "Привет! Этот бот поможет зарегистрироваться на кабинетные ролевые игры.\n\n"
            "Сначала заполните профиль и запишитесь на игру. Мы сообщим вам, когда будет открыто "
            "подтверждение на игру — вы сможете подтвердиться, когда будете уверены, что сможете прийти на игру."
        )
        assert [button.value for button in response.buttons] == [
            PROFILE,
            ENLIST,
            CONFIRM,
            GAME_WISHES,
            CANCEL,
            MY_GAMES,
        ]
        assert saved is not None and saved.dialog_state == "IDLE"
        assert saved.dialog_context == {}

    assert not publisher.commands


@pytest.mark.asyncio
async def test_vk_start_message_matches_slash_start_greeting(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, users, publisher, _ = await engine_setup(disk_store, event)

    def vk(update: int, value: str) -> InboundMessage:
        return InboundMessage(
            identity=BotIdentity(platform=Platform.VK, platform_user_id=7),
            update_id=str(update),
            text=value,
            peer_id=7,
        )

    slash_start = await engine.handle(vk(1, "/start"))
    await engine.handle(vk(2, PROFILE))
    vk_start = await engine.handle(vk(3, "Начать"))
    saved = await users.get(Platform.VK, 7)

    assert vk_start == slash_start
    assert saved is not None and saved.dialog_state == "IDLE"
    assert saved.dialog_context == {}
    assert not publisher.commands


@pytest.mark.asyncio
async def test_every_free_text_state_rejects_a_button_absent_from_the_last_bot_message(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, publisher, _ = await engine_setup(disk_store, event, admin=True)
    current_button = Button(label="Текущая кнопка", value="current:button")

    for update, state in enumerate(sorted(FREE_TEXT_DIALOG_STATES), start=1):
        user = await users.get(Platform.TELEGRAM, 1)
        assert user is not None
        user.dialog_state = state
        user.dialog_context = {"sentinel": state}
        user.last_bot_buttons = [current_button]
        await users.save(user)

        response = await engine.handle(inbound(update, "stale:button", callback=True))
        saved = await users.get(Platform.TELEGRAM, 1)

        assert "кнопка устарела" in response.text.casefold()
        assert response.buttons == [current_button]
        assert saved is not None and saved.dialog_state == state
        assert saved.dialog_context == {"sentinel": state}

    assert not publisher.commands


@pytest.mark.asyncio
async def test_created_game_accepts_signup_but_does_not_offer_confirmation(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    event.status = EventStatus.CREATED
    engine, _, publisher, tables = await engine_setup(disk_store, event)

    enlist_games = await engine.handle(inbound(1, ENLIST))
    assert any(button.value == f"select:enlist:{event.event_id}" for button in enlist_games.buttons)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Алиса",
    )

    confirmation_games = await engine.handle(inbound(2, CONFIRM))
    assert confirmation_games.text == "Сейчас нет игр, доступных для подтверждения участия."
    assert not any(button.value == f"select:confirm:{event.event_id}" for button in confirmation_games.buttons)
    assert not publisher.commands


@pytest.mark.asyncio
async def test_confirmation_picker_filters_games_before_paginating(disk_store: MemoryDiskStore) -> None:
    base = datetime(2026, 1, 1, tzinfo=UTC)
    game_date = date(2026, 11, 14)
    events = [
        Event(
            event_id=f"event-{index:04d}",
            name=(
                f"Unavailable {index}"
                if index < PAGE_SIZE
                else f"Eligible {index - PAGE_SIZE}"
                if index < PAGE_SIZE + 3
                else "Already confirmed"
            ),
            game_date=game_date,
            disk_resource_path=f"disk:/larp-bot/events/event-{index:04d}/game.xlsx",
            public_registration_url=f"https://disk.example/{index}",
            status=EventStatus.CREATED if index < PAGE_SIZE else EventStatus.CONFIRMATION_OPEN,
            registrations_migrated_at=base,
            created_at=base + timedelta(seconds=index),
            updated_at=base + timedelta(seconds=index),
        )
        for index in range(PAGE_SIZE + 4)
    ]
    engine, _, _, registrations = await engine_setup(disk_store, *events)
    await registrations.import_missing(
        [
            Registration(
                event_id=candidate.event_id,
                participant_key=engine.registrations.key(Platform.TELEGRAM, 1, candidate.event_id),
                display_name="Иван Иванов",
                wish_play="Без пожеланий",
                attendance_status=(
                    AttendanceStatus.CANCELLED
                    if candidate is events[PAGE_SIZE - 1]
                    else AttendanceStatus.CONFIRMED
                    if candidate is events[-1]
                    else AttendanceStatus.WAITING
                ),
            )
            for candidate in events
        ]
    )

    response = await engine.handle(inbound(1, CONFIRM))
    game_buttons = [button for button in response.buttons if button.value.startswith("select:confirm:")]

    assert response.text == "✅ Выберите игру, участие в которой хотите подтвердить:"
    assert [button.label for button in game_buttons] == ["Eligible 0", "Eligible 1", "Eligible 2"]
    assert PAGE_FORWARD not in {button.label for button in response.buttons}
    assert "Already confirmed" not in {button.label for button in response.buttons}


@pytest.mark.asyncio
async def test_game_archived_during_signup_is_not_enqueued(disk_store: MemoryDiskStore, event: Event) -> None:
    event.status = EventStatus.CREATED
    engine, _, publisher, _ = await engine_setup(disk_store, event)

    games = await engine.handle(inbound(1, ENLIST))
    assert any(button.value == f"select:enlist:{event.event_id}" for button in games.buttons)
    await engine.handle(inbound(2, f"select:enlist:{event.event_id}", callback=True))
    await engine.events.archive(event.event_id, datetime.now(UTC))

    rejected = await engine.handle(inbound(3, "enlist:wish-play:skip", callback=True))
    hidden = await engine.handle(inbound(4, ENLIST))
    stale = await engine.handle(inbound(5, f"select:enlist:{event.event_id}", callback=True))

    assert rejected.text == "Игра больше не активна; запись и изменение данных недоступны."
    assert not any(button.value == f"select:enlist:{event.event_id}" for button in hidden.buttons)
    assert stale.text == "Игра больше не активна; запись и изменение данных недоступны."
    assert publisher.commands == []


@pytest.mark.asyncio
async def test_game_archived_during_cancellation_is_not_enqueued(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, publisher, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Алиса",
    )

    games = await engine.handle(inbound(1, CANCEL))
    assert any(button.value == f"select:cancel:{event.event_id}" for button in games.buttons)
    await engine.handle(inbound(2, f"select:cancel:{event.event_id}", callback=True))
    await engine.events.archive(event.event_id, datetime.now(UTC))

    rejected = await engine.handle(inbound(3, "cancel:yes", callback=True))
    hidden = await engine.handle(inbound(4, CANCEL))
    stale = await engine.handle(inbound(5, f"select:cancel:{event.event_id}", callback=True))

    assert rejected.text == "Игра больше не активна; запись и изменение данных недоступны."
    assert not any(button.value == f"select:cancel:{event.event_id}" for button in hidden.buttons)
    assert hidden.text == "Сейчас нет игр, доступных для отмены участия."
    assert hidden.buttons == [Button(label=MAIN_MENU_LABEL, value="Отмена")]
    assert stale.text == "Игра больше не активна; запись и изменение данных недоступны."
    assert publisher.commands == []


@pytest.mark.asyncio
async def test_closed_game_still_offers_and_accepts_cancellation(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, publisher, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Алиса",
    )

    games = await engine.handle(inbound(1, CANCEL))
    assert any(button.value == f"select:cancel:{event.event_id}" for button in games.buttons)
    await engine.handle(inbound(2, f"select:cancel:{event.event_id}", callback=True))
    await engine.events.set_status(event.event_id, EventStatus.CLOSED)

    accepted = await engine.handle(inbound(3, "cancel:yes", callback=True))
    still_visible = await engine.handle(inbound(4, CANCEL))
    selected_again = await engine.handle(inbound(5, f"select:cancel:{event.event_id}", callback=True))

    assert accepted.text == "⏳ Отмена принята в обработку."
    assert accepted.deferred and accepted.command_enqueued
    assert any(button.value == f"select:cancel:{event.event_id}" for button in still_visible.buttons)
    assert "Вы уверены, что хотите отменить участие" in selected_again.text
    assert [command.operation for command in publisher.commands] == [Operation.CANCEL]


@pytest.mark.asyncio
async def test_cancellation_picker_shows_only_registered_closed_game_without_pagination(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    base = datetime(2026, 1, 1, tzinfo=UTC)
    event.status = EventStatus.CLOSED
    event.registrations_migrated_at = base
    event.created_at = base
    unrelated_open_events = [
        event.model_copy(
            update={
                "event_id": f"unrelated-{index}",
                "name": f"Unrelated {index}",
                "status": EventStatus.CREATED,
                "registrations_migrated_at": base,
                "created_at": base + timedelta(seconds=index + 1),
                "updated_at": base + timedelta(seconds=index + 1),
            }
        )
        for index in range(PAGE_SIZE + 1)
    ]
    engine, _, _, registrations = await engine_setup(disk_store, event, *unrelated_open_events)
    await registrations.import_missing(
        [
            Registration(
                event_id=event.event_id,
                participant_key=engine.registrations.key(Platform.TELEGRAM, 1, event.event_id),
                display_name="Иван Иванов",
                wish_play="Без пожеланий",
            )
        ]
    )

    response = await engine.handle(inbound(1, CANCEL))

    assert response.text == "❌ Выберите игру, участие в которой хотите отменить:"
    assert response.buttons == [
        Button(label=event.name, value=f"select:cancel:{event.event_id}"),
        Button(label=MAIN_MENU_LABEL, value="Отмена"),
    ]


@pytest.mark.asyncio
async def test_archived_game_does_not_offer_or_open_character_wish_editing(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, _, publisher, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Алиса",
    )
    await tables.confirm(
        event.event_id,
        operation_id="existing-confirm",
        participant_key=key,
        character_wish="Doctor",
    )
    await engine.events.archive(event.event_id, datetime.now(UTC))

    games = await engine.handle(inbound(1, GAME_WISHES))
    stale = await engine.handle(inbound(2, f"select:character:{event.event_id}", callback=True))

    assert not any(button.value == f"select:character:{event.event_id}" for button in games.buttons)
    assert games.text == "Нет игр, на которые вы записаны."
    assert stale.text == "Игра больше не активна; запись и изменение данных недоступны."
    assert publisher.commands == []


@pytest.mark.asyncio
async def test_game_archived_during_character_wish_edit_is_not_enqueued(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, _, publisher, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван Иванов",
        wish_play="Алиса",
    )
    await tables.confirm(
        event.event_id,
        operation_id="existing-confirm",
        participant_key=key,
        character_wish="Doctor",
    )

    await engine.handle(inbound(1, GAME_WISHES))
    await engine.handle(inbound(2, f"wishes:character:{event.event_id}", callback=True))
    await engine.events.archive(event.event_id, datetime.now(UTC))

    rejected = await engine.handle(inbound(3, "Medic"))

    assert rejected.text == "Игра больше не активна; запись и изменение данных недоступны."
    assert publisher.commands == []


@pytest.mark.asyncio
async def test_cancelled_registration_must_reenlist_before_confirmation(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, _, publisher, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван",
        wish_play="A",
    )
    await tables.confirm(event.event_id, operation_id="confirm", participant_key=key, character_wish="Doctor")
    await tables.cancel(event.event_id, operation_id="cancel", participant_key=key)
    listing = await engine.handle(inbound(1, CONFIRM))
    selected = await engine.handle(inbound(2, f"select:confirm:{event.event_id}", callback=True))

    assert not any(button.value == f"select:confirm:{event.event_id}" for button in listing.buttons)
    assert "сначала запишитесь на игру снова" in listing.text.casefold()
    assert selected.text == "Сначала снова запишитесь на эту игру."
    assert publisher.commands == []


@pytest.mark.asyncio
async def test_vk_uses_same_profile_and_registration_engine(disk_store: MemoryDiskStore, event: Event) -> None:
    users = MemoryUserRepository()
    events = MemoryEventRepository([event])
    tables = MemoryRegistrationRepository()
    showcase = YandexDiskShowcaseRepository(disk_store)
    await showcase.create_event_workbook(event.disk_resource_path)
    catalog = RegistrationCatalog(events, tables, showcase)
    publisher = MemoryCommandPublisher()
    engine = ConversationEngine(
        users,
        events,
        RegistrationService(events, catalog, publisher, "secret"),
        EventAdministrationService(events, showcase, catalog, users, "secret"),
        StaticAdminProvider(vk_ids={7}),
    )

    def vk(update: int, value: str, *, callback: bool = False) -> InboundMessage:
        return InboundMessage(
            identity=BotIdentity(platform=Platform.VK, platform_user_id=7),
            update_id=str(update),
            text="" if callback else value,
            callback=value if callback else None,
            peer_id=7,
        )

    for update, value in enumerate((PROFILE, "Иванов Иван", "Нет", "Пропустить", "Да"), start=1):
        await engine.handle(vk(update, value))
    profile = await users.get(Platform.VK, 7)
    assert isinstance(profile, VkUser)
    assert profile.profile_complete and profile.telegram_handle is None
    assert profile.full_name == "Иванов Иван"

    await engine.handle(vk(10, ENLIST))
    await engine.handle(vk(11, f"select:enlist:{event.event_id}", callback=True))
    await engine.handle(vk(12, "Алиса"))
    enlist = publisher.commands[-1]
    assert enlist.operation is Operation.ENLIST
    assert isinstance(enlist.payload, EnlistPayload)
    assert enlist.participant_key is not None
    assert enlist.payload.vk_profile == "https://vk.com/id7"
    assert enlist.payload.telegram_profile is None
    await tables.enlist(
        event.event_id,
        operation_id=enlist.operation_id,
        participant_key=enlist.participant_key,
        display_name=enlist.payload.display_name,
        wish_play=enlist.payload.wish_play,
        larp_experience=enlist.payload.larp_experience,
        crossplay=enlist.payload.crossplay,
    )
    await engine.handle(vk(14, CONFIRM))
    prompt = await engine.handle(vk(15, f"select:confirm:{event.event_id}", callback=True))
    assert "пожелания" in prompt.text.casefold()
    stale = await engine.handle(vk(16, f"select:confirm:{event.event_id}", callback=True))
    assert "кнопка устарела" in stale.text.casefold()
    assert stale.buttons == prompt.buttons
    assert [command.operation for command in publisher.commands] == [Operation.SYNC_PROFILE, Operation.ENLIST]
    await engine.handle(vk(17, "Doctor"))
    confirmation = publisher.commands[-1]
    assert confirmation.operation is Operation.CONFIRM
    assert confirmation.participant_key is not None
    assert hasattr(confirmation.payload, "character_wish")
    await tables.confirm(
        event.event_id,
        operation_id=confirmation.operation_id,
        participant_key=confirmation.participant_key,
        character_wish=confirmation.payload.character_wish,
    )

    await engine.handle(vk(18, GAME_WISHES))
    edit_prompt = await engine.handle(vk(19, f"wishes:character:{event.event_id}", callback=True))
    assert "Отправьте новый вариант" in edit_prompt.text
    await engine.handle(vk(20, "Medic"))
    assert publisher.commands[-1].operation is Operation.UPDATE_CHARACTER_WISH

    menu = await engine.handle(vk(21, ADMIN))
    assert "Администрирование" in menu.text


@pytest.mark.parametrize(
    ("telegram_profile", "expected_handle"),
    [
        ("@moscowkudesnik", "@moscowkudesnik"),
        ("moscowkudesnik", "@moscowkudesnik"),
        ("t.me/moscowkudesnik", "@moscowkudesnik"),
        ("https://t.me/moscowkudesnik", "@moscowkudesnik"),
    ],
)
@pytest.mark.asyncio
async def test_vk_profile_parses_telegram_reference(
    disk_store: MemoryDiskStore,
    event: Event,
    telegram_profile: str,
    expected_handle: str | None,
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event)

    def vk(update: int, value: str) -> InboundMessage:
        return InboundMessage(
            identity=BotIdentity(platform=Platform.VK, platform_user_id=7),
            update_id=str(update),
            text=value,
            peer_id=7,
        )

    responses = []
    for update, value in enumerate(
        (PROFILE, "Иванов Иван", "Нет", telegram_profile, "Да"),
        start=1,
    ):
        responses.append(await engine.handle(vk(update, value)))

    profile = await users.get(Platform.VK, 7)
    assert responses[3].text == "Готовы ли вы играть за персонажа противоположного пола?"
    assert isinstance(profile, VkUser)
    assert profile.profile_complete
    assert profile.telegram_handle == expected_handle


@pytest.mark.asyncio
async def test_vk_profile_retries_malformed_telegram_reference_with_skip_option(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event)

    def vk(update: int, value: str, *, callback: bool = False) -> InboundMessage:
        return InboundMessage(
            identity=BotIdentity(platform=Platform.VK, platform_user_id=7),
            update_id=str(update),
            text="" if callback else value,
            callback=value if callback else None,
            peer_id=7,
        )

    await engine.handle(vk(1, PROFILE))
    await engine.handle(vk(2, "Иванов Иван"))
    await engine.handle(vk(3, "Нет"))
    rejected = await engine.handle(vk(4, "not a telegram profile"))
    pending = await users.get(Platform.VK, 7)

    assert rejected.text == "Не удалось разобрать профиль Telegram. Попробуйте ещё раз"
    assert [button.label for button in rejected.buttons] == ["Пропустить", MAIN_MENU_LABEL]
    assert pending is not None and pending.dialog_state == "PROFILE_CONTACT"

    await engine.handle(vk(5, "Пропустить", callback=True))
    await engine.handle(vk(6, "Да"))
    profile = await users.get(Platform.VK, 7)
    assert isinstance(profile, VkUser) and profile.profile_complete
    assert profile.telegram_handle is None


@pytest.mark.asyncio
async def test_telegram_identity_handle_is_refreshed_from_each_update(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event)

    await engine.handle(inbound(1, "/start", user_id=2, telegram_username="Current_Name"))

    user = await users.get(Platform.TELEGRAM, 2)
    assert isinstance(user, TelegramUser)
    assert user.telegram_handle == "@current_name"


@pytest.mark.asyncio
async def test_waiting_blank_character_menu_routes_to_confirmation(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, _, tables = await engine_setup(disk_store, event)
    key = engine.registrations.key(Platform.TELEGRAM, 1, event.event_id)
    await tables.enlist(
        event.event_id,
        operation_id="existing-enlist",
        participant_key=key,
        display_name="Иван",
        wish_play="A",
    )
    await engine.handle(inbound(1, GAME_WISHES))
    response = await engine.handle(inbound(2, f"wishes:character:{event.event_id}", callback=True))
    assert CONFIRM in response.text


@pytest.mark.asyncio
async def test_incomplete_profile_cannot_enlist(disk_store: MemoryDiskStore, event: Event) -> None:
    users = MemoryUserRepository()
    events = MemoryEventRepository([event])
    tables = MemoryRegistrationRepository()
    showcase = YandexDiskShowcaseRepository(disk_store)
    catalog = RegistrationCatalog(events, tables, showcase)
    publisher = MemoryCommandPublisher()
    engine = ConversationEngine(
        users,
        events,
        RegistrationService(events, catalog, publisher, "secret"),
        EventAdministrationService(events, showcase, catalog, users, "secret"),
        StaticAdminProvider(),
    )
    response = await engine.handle(inbound(1, ENLIST))
    assert "Сначала заполните профиль" in response.text
    assert response.buttons[0].value == PROFILE

    forged = await engine.handle(inbound(2, f"select:enlist:{event.event_id}", callback=True))
    assert "полностью заполните профиль" in forged.text
    assert not publisher.commands


@pytest.mark.asyncio
async def test_profile_validates_email_before_advancing_to_next_question(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event)
    user_id = 2
    answers = (
        PROFILE,
        "Иванов Иван Иванович",
        "Да",
        "Да",
        "+7 999 123-45-67",
    )
    for update, answer in enumerate(answers, start=1):
        await engine.handle(inbound(update, answer, user_id=user_id))

    invalid = await engine.handle(inbound(6, "not-an-email", user_id=user_id))

    assert "Некорректный email" in invalid.text
    pending = await users.get(Platform.TELEGRAM, user_id)
    assert pending is not None
    assert pending.dialog_state == "PROFILE_PASS_EMAIL_ADDRESS"
    assert "email" not in pending.dialog_context

    contact_prompt = await engine.handle(inbound(7, "player@example.com", user_id=user_id))
    await engine.handle(inbound(8, "https://vk.com/id2", user_id=user_id))
    saved = await engine.handle(inbound(9, "Да", user_id=user_id))
    profile = await users.get(Platform.TELEGRAM, user_id)
    assert "страницу VK" in contact_prompt.text
    assert "Профиль успешно заполнен" in saved.text
    assert profile is not None and profile.pass_details is not None
    assert profile.pass_details.email == "player@example.com"
    assert profile.pass_details.mobile_phone == "+7 999 123-45-67"
    assert profile.pass_details.foreigner is False
    assert profile.pass_details.surname_latin is None
    assert profile.full_name == "Иванов Иван"


@pytest.mark.asyncio
async def test_profile_collects_one_cyrillic_full_name_and_keeps_hyphenated_parts(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event)

    full_name_prompt = await engine.handle(inbound(1, PROFILE, user_id=2))
    invalid_full_name = await engine.handle(inbound(2, "Smith John", user_id=2))
    pass_prompt = await engine.handle(inbound(3, "  Семёнова-Звенигородская   Ева-Лотта Ильинична ", user_id=2))
    contact_prompt = await engine.handle(inbound(4, "Нет", user_id=2))
    for update, answer in enumerate(("https://vk.com/id2", "Нет"), start=5):
        saved = await engine.handle(inbound(update, answer, user_id=2))

    profile = await users.get(Platform.TELEGRAM, 2)
    assert full_name_prompt.text == "Введите ФИО кириллицей (фамилию имя, если отчества нету):"
    assert [button.value for button in full_name_prompt.buttons] == ["Отмена"]
    assert invalid_full_name.text == "Введите ФИО кириллицей, либо фамилию имя, если отчества нету"
    assert [button.value for button in invalid_full_name.buttons] == ["Отмена"]
    assert pass_prompt.text == f"Нужен ли вам пропуск для входа на место проведения игры ({VENUE_ADDRESS})?"
    assert "страницу VK" in contact_prompt.text
    assert "Профиль успешно заполнен" in saved.text
    assert profile is not None
    assert profile.full_name == "Семёнова-Звенигородская Ева-Лотта"
    assert profile.larp_experience is None
    assert profile.needs_pass is False and profile.pass_details is None


@pytest.mark.asyncio
async def test_profile_completion_shows_confirmation_before_main_menu(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, _, _, _ = await engine_setup(disk_store, event)
    responses = []
    for update, answer in enumerate(
        (PROFILE, "Иванов Иван", "Нет", "https://vk.com/id2", "Нет"),
        start=1,
    ):
        responses.append(await engine.handle(inbound(update, answer, user_id=2)))

    completed = responses[-1]
    assert completed.text == "✅ Профиль успешно заполнен."
    assert completed.buttons == [Button(label=MAIN_MENU_LABEL, value=MAIN_MENU)]
    assert not completed.deferred and not completed.command_enqueued

    menu = await engine.handle(inbound(6, MAIN_MENU, callback=True, user_id=2))
    assert [button.value for button in menu.buttons] == [PROFILE, ENLIST, CONFIRM, GAME_WISHES, CANCEL, MY_GAMES]


@pytest.mark.asyncio
async def test_profile_completion_queues_experience_lookup_without_asking_the_player(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, users, publisher, _ = await engine_setup(disk_store, event)
    responses = []
    for update, answer in enumerate(
        (PROFILE, "Иванов Иван", "Нет", "https://vk.com/id2", "Нет"),
        start=1,
    ):
        responses.append(await engine.handle(inbound(update, answer, user_id=2)))

    profile = await users.get(Platform.TELEGRAM, 2)
    shown = await engine.handle(inbound(6, PROFILE, user_id=2))

    assert "Профиль успешно заполнен" in responses[-1].text
    assert all("Играли ли вы" not in response.text for response in responses)
    assert profile is not None and profile.larp_experience is None
    assert len(publisher.commands) == 1
    sync = publisher.commands[0]
    assert sync.operation is Operation.SYNC_PROFILE
    assert sync.event_id == "profile:telegram:2"
    assert sync.payload == ProfileSyncPayload(names=["Иванов Иван"])
    assert "Опыт LARP" not in shown.text
    assert "Кросспол:" in shown.text
    assert "Кроссплей:" not in shown.text


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("admin", "gamemaster", "expected_status"),
    [
        (False, False, "Игрок"),
        (False, True, "Модератор"),
        (True, True, "Администратор"),
    ],
)
async def test_profile_shows_bot_user_status_with_admin_precedence(
    disk_store: MemoryDiskStore,
    event: Event,
    *,
    admin: bool,
    gamemaster: bool,
    expected_status: str,
) -> None:
    engine, _, _, _ = await engine_setup(
        disk_store,
        event,
        admin=admin,
        gamemaster=gamemaster,
    )

    shown = await engine.handle(inbound(1, PROFILE))

    assert f"Статус: {expected_status}" in shown.text


@pytest.mark.asyncio
async def test_profile_shows_persisted_moderator_status(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event)
    moderator = await users.get(Platform.TELEGRAM, 1)
    assert moderator is not None
    moderator.is_gamemaster = True
    await users.save(moderator)

    shown = await engine.handle(inbound(1, PROFILE))

    assert "Статус: Модератор" in shown.text


@pytest.mark.asyncio
async def test_profile_save_queues_one_background_sync_instead_of_scanning_games(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    second = event.model_copy(
        update={
            "event_id": "event-second",
            "disk_resource_path": "disk:/larp-bot/events/event-second/master_table_second.xlsx",
            "public_registration_url": "https://disk.example/second",
            "created_at": event.created_at + timedelta(seconds=1),
        }
    )
    archived = event.model_copy(
        update={
            "event_id": "event-archived",
            "disk_resource_path": "disk:/larp-bot/events/event-archived/master_table_archived.xlsx",
            "public_registration_url": "https://disk.example/archived",
            "archived_at": datetime.now(UTC),
            "created_at": event.created_at + timedelta(seconds=2),
        }
    )
    leader_only = event.model_copy(
        update={
            "event_id": "event-led",
            "disk_resource_path": "disk:/larp-bot/events/event-led/master_table_led.xlsx",
            "public_registration_url": "https://disk.example/led",
            "created_at": event.created_at + timedelta(seconds=3),
        }
    )
    engine, users, publisher, tables = await engine_setup(disk_store, event, second, archived, leader_only)
    for game in (event, second, archived):
        await tables.enlist(
            game.event_id,
            operation_id=f"enlist-{game.event_id}",
            participant_key=engine.registrations.key(Platform.TELEGRAM, 1, game.event_id),
            display_name="Иван Иванов",
            wish_play="Без пожеланий",
        )
    await engine.events.add_leader(
        leader_only.event_id,
        BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1),
    )

    answers = (
        "profile:edit",
        "Новиков Новый Нович",
        "Нет",
        "https://vk.com/new-profile",
        "Нет",
    )
    for update, answer in enumerate(answers, start=1):
        response = await engine.handle(inbound(update, answer, user_id=1, telegram_username="new_profile"))

    profile = await users.get(Platform.TELEGRAM, 1)
    assert "Профиль успешно заполнен" in response.text
    assert profile is not None and profile.full_name == "Новиков Новый"
    assert profile.larp_experience is True
    assert profile.last_update_id == "5"
    assert len(publisher.commands) == 1
    sync = publisher.commands[0]
    assert sync.operation is Operation.SYNC_PROFILE
    assert sync.event_id == "profile:telegram:1"
    assert sync.payload == ProfileSyncPayload(names=["Иван Иванов", "Новиков Новый"])


@pytest.mark.asyncio
async def test_foreign_profile_collects_combined_latin_full_name(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event)
    answers = (
        PROFILE,
        "Ли Анна",
        "Да",
        "Нет",
        "Li Анна",
        "Li Anna",
        "+44 7700 900123",
        "anna@example.com",
        "https://vk.com/anna-li",
        "Нет",
    )
    responses = []
    for update, answer in enumerate(answers, start=1):
        responses.append(await engine.handle(inbound(update, answer, user_id=2)))

    profile = await users.get(Platform.TELEGRAM, 2)
    assert responses[3].text == "Введите ФИО латиницей (фамилию имя, если отчества нету):"
    assert responses[4].text == "Введите ФИО латиницей, либо фамилию имя, если отчества нету"
    assert [button.value for button in responses[4].buttons] == ["Отмена"]
    assert "Профиль успешно заполнен" in responses[-1].text
    assert profile is not None and profile.profile_complete and profile.pass_details is not None
    assert profile.pass_details.model_dump() == {
        "surname_cyrillic": "Ли",
        "name_cyrillic": "Анна",
        "patronym_cyrillic": "-",
        "foreigner": True,
        "surname_latin": "Li",
        "name_latin": "Anna",
        "patronym_latin": "-",
        "mobile_phone": "+44 7700 900123",
        "email": "anna@example.com",
    }


@pytest.mark.asyncio
async def test_admin_archive_requires_exact_case_but_trims_whitespace(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, _, publisher, _ = await engine_setup(disk_store, event, admin=True)
    menu = await engine.handle(inbound(1, ADMIN))
    assert MANAGE_GAMES in [button.value for button in menu.buttons]
    await engine.handle(inbound(2, MANAGE_GAMES, callback=True))
    management = await engine.handle(inbound(3, f"ag:manage:{event.event_id}", callback=True))
    assert ARCHIVE_GAME in [button.label for button in management.buttons]
    await engine.handle(inbound(4, f"ag:archive:{event.event_id}", callback=True))
    wrong = await engine.handle(inbound(5, event.name.casefold()))
    assert "не совпало" in wrong.text
    accepted = await engine.handle(inbound(6, f"  {event.name}  "))
    assert accepted.deferred
    assert [(button.label, button.value) for button in accepted.buttons] == [("⬅️ Назад", MANAGE_GAMES)]
    assert [(button.label, button.value) for button in publisher.commands[-1].reply_context.buttons] == [
        ("⬅️ Назад", MANAGE_GAMES)
    ]
    assert publisher.commands[-1].operation is Operation.ARCHIVE_EVENT


@pytest.mark.asyncio
async def test_gamemaster_has_privileged_interface_but_cannot_mutate_an_unled_game(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    event = Event.model_validate(
        event.model_dump()
        | {
            "pass_table_resource_path": f"disk:/larp-bot/passes/{event.event_id}.xlsx",
            "pass_table_public_url": "https://disk.example/pass-table",
        }
    )
    engine, _, publisher, _ = await engine_setup(disk_store, event, gamemaster=True)

    main = await engine.handle(inbound(1, "/start"))
    assert ADMIN in [button.value for button in main.buttons]

    menu = await engine.handle(inbound(2, ADMIN))
    assert [button.value for button in menu.buttons] == [
        "➕ Создать игру",
        MANAGE_GAMES,
        "⬅️ Назад",
    ]

    games = await engine.handle(inbound(3, MANAGE_GAMES))
    assert any(button.value == f"ag:manage:{event.event_id}" for button in games.buttons)
    management = await engine.handle(inbound(4, f"ag:manage:{event.event_id}", callback=True))
    assert "доступны только мастерам" in management.text
    assert EVENT_TABLES not in [button.label for button in management.buttons]
    assert "https://disk.example/pass-table" not in management.text

    tables_denied = await engine.handle(inbound(5, f"ag:tables:{event.event_id}", callback=True))
    assert tables_denied.text == "Только мастера этой игры могут изменять её данные."
    assert "https://disk.example/pass-table" not in tables_denied.text
    notification_denied = await engine.handle(inbound(6, f"ag:notify:{event.event_id}", callback=True))
    assert notification_denied.text == "Только мастера этой игры могут изменять её данные."
    waiting_denied = await engine.handle(inbound(7, f"ag:waiting:{event.event_id}", callback=True))
    assert waiting_denied.text == "Только мастера этой игры могут изменять её данные."
    removal_denied = await engine.handle(inbound(8, f"ag:remove:{event.event_id}", callback=True))
    assert removal_denied.text == "Только мастера этой игры могут изменять её данные."
    assert publisher.commands == []


@pytest.mark.asyncio
async def test_event_leader_pages_waiting_players_and_shows_signup_profiles(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, users, _, tables = await engine_setup(disk_store, event, admin=True)
    base = datetime(2026, 1, 1, tzinfo=UTC)
    registrations: list[Registration] = []
    for index in range(10):
        player = TelegramUser(
            tg_id=100 + index,
            full_name=f"Игрок Telegram {index}",
            vk_url=f"https://vk.com/id{100 + index}",
            telegram_handle=f"@telegram_player_{index}",
            crossplay=True,
            larp_experience=True,
            needs_pass=False,
        )
        await users.save(player)
        registrations.append(
            Registration(
                event_id=event.event_id,
                participant_key=engine.registrations.key(Platform.TELEGRAM, player.tg_id, event.event_id),
                display_name=player.full_name or "",
                vk_profile=player.vk_url or "",
                telegram_profile=f"https://t.me/{player.telegram_handle.removeprefix('@')}",
                wish_play="Без пожеланий",
                created_at=base + timedelta(seconds=index),
            )
        )

    vk_player = VkUser(
        vk_id=250,
        full_name="Игрок ВК",
        telegram_handle="@vk_player_profile",
        crossplay=False,
        larp_experience=True,
        needs_pass=False,
    )
    await users.save(vk_player)
    registrations.append(
        Registration(
            event_id=event.event_id,
            participant_key=engine.registrations.key(Platform.VK, vk_player.vk_id, event.event_id),
            display_name=vk_player.full_name or "",
            vk_profile="https://vk.com/id250",
            telegram_profile="https://t.me/vk_player_profile",
            wish_play="Без пожеланий",
            created_at=base + timedelta(seconds=10),
        )
    )
    await tables.import_missing(
        [
            *registrations,
            Registration(
                event_id=event.event_id,
                participant_key="c" * 43,
                display_name="Подтвердивший игрок",
                wish_play="Без пожеланий",
                attendance_status=AttendanceStatus.CONFIRMED,
                created_at=base + timedelta(seconds=4, microseconds=1),
            ),
            Registration(
                event_id=event.event_id,
                participant_key="x" * 43,
                display_name="Отменивший игрок",
                wish_play="Без пожеланий",
                attendance_status=AttendanceStatus.CANCELLED,
                created_at=base + timedelta(seconds=5, microseconds=1),
            ),
        ]
    )

    management = await engine.handle(inbound(1, f"ag:manage:{event.event_id}", callback=True))
    assert LIST_WAITING_PLAYERS in [button.label for button in management.buttons]

    first = await engine.handle(inbound(2, f"ag:waiting:{event.event_id}", callback=True))
    first_players = [button for button in first.buttons if button.value.startswith("admin:waiting:pick:")]
    assert len(first_players) == PAGE_SIZE
    assert [button.label for button in first_players] == [f"Игрок Telegram {index}" for index in range(PAGE_SIZE)]
    assert "Подтвердивший игрок" not in [button.label for button in first_players]
    assert "Отменивший игрок" not in [button.label for button in first_players]

    telegram_profile = await engine.handle(inbound(3, first_players[0].value, callback=True))
    assert telegram_profile.text == (
        "Этот пользователь зарегистрировался через бота Telegram. "
        "Вот его профиль вк: https://vk.com/id100\n"
        "Вот его профиль Telegram: https://t.me/telegram_player_0"
    )

    returned = await engine.handle(inbound(4, f"ag:waiting:{event.event_id}", callback=True))
    next_button = next(button for button in returned.buttons if button.label == PAGE_FORWARD)
    second = await engine.handle(inbound(5, next_button.value, callback=True))
    second_players = [button for button in second.buttons if button.value.startswith("admin:waiting:pick:")]
    assert [button.label for button in second_players] == [f"Игрок Telegram {index}" for index in range(5, 10)]

    next_button = next(button for button in second.buttons if button.label == PAGE_FORWARD)
    third = await engine.handle(inbound(6, next_button.value, callback=True))
    third_players = [button for button in third.buttons if button.value.startswith("admin:waiting:pick:")]
    assert [button.label for button in third_players] == ["Игрок ВК"]

    vk_profile = await engine.handle(inbound(7, third_players[0].value, callback=True))
    assert vk_profile.text == (
        "Этот пользователь зарегистрировался через бота ВК. "
        "Вот его профиль вк: https://vk.com/id250\n"
        "Вот его профиль Telegram: https://t.me/vk_player_profile"
    )


@pytest.mark.asyncio
async def test_event_leader_pages_profile_names_and_confirms_player_removal(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, _, publisher, tables = await engine_setup(disk_store, event, admin=True)
    base = datetime(2026, 1, 1, tzinfo=UTC)
    registrations = [
        Registration(
            event_id=event.event_id,
            participant_key=f"{index:043d}",
            display_name=f"Фамилия{index} Имя{index}",
            wish_play="Без пожеланий",
            attendance_status=(AttendanceStatus.CONFIRMED if index % 2 else AttendanceStatus.WAITING),
            created_at=base + timedelta(seconds=index),
        )
        for index in range(11)
    ]
    await tables.import_missing(
        [
            *registrations,
            Registration(
                event_id=event.event_id,
                participant_key="c" * 43,
                display_name="Отменённый Игрок",
                wish_play="Без пожеланий",
                attendance_status=AttendanceStatus.CANCELLED,
                created_at=base + timedelta(seconds=5, microseconds=1),
            ),
        ]
    )

    management = await engine.handle(inbound(1, f"ag:manage:{event.event_id}", callback=True))
    assert MANAGE_PLAYERS in [button.label for button in management.buttons]
    assert REMOVE_PLAYER not in [button.label for button in management.buttons]

    player_management = await engine.handle(inbound(2, f"ag:players:{event.event_id}", callback=True))
    assert [button.label for button in player_management.buttons[:2]] == [REMOVE_PLAYER, REMOVE_WAITING_PLAYERS]

    first = await engine.handle(inbound(3, f"ag:remove:{event.event_id}", callback=True))
    first_players = [button for button in first.buttons if button.value.startswith("admin:remove:pick:")]
    assert len(first_players) == PAGE_SIZE
    assert first_players[0].label == "Фамилия0 Имя0 — Ожидает подтверждения"
    assert "Отменённый Игрок" not in [button.label for button in first_players]
    assert "имена из их профилей" in first.text

    next_button = next(button for button in first.buttons if button.label == PAGE_FORWARD)
    second = await engine.handle(inbound(4, next_button.value, callback=True))
    second_players = [button for button in second.buttons if button.value.startswith("admin:remove:pick:")]
    assert [button.label for button in second_players] == [
        "Фамилия5 Имя5 — Подтверждено",
        "Фамилия6 Имя6 — Ожидает подтверждения",
        "Фамилия7 Имя7 — Подтверждено",
        "Фамилия8 Имя8 — Ожидает подтверждения",
        "Фамилия9 Имя9 — Подтверждено",
    ]

    stale = await engine.handle(inbound(5, first_players[0].value, callback=True))
    assert "кнопка устарела" in stale.text.casefold()
    assert publisher.commands == []

    previous_button = next(button for button in second.buttons if button.label == PAGE_BACK)
    returned = await engine.handle(inbound(6, previous_button.value, callback=True))
    target = next(button for button in returned.buttons if button.label.startswith("Фамилия1 Имя1 —"))
    warning = await engine.handle(inbound(7, target.value, callback=True))

    assert "Это действие нельзя отменить" in warning.text
    assert "Вы уверены, что выбрали нужного человека?" in warning.text
    assert "Имя из профиля: Фамилия1 Имя1" in warning.text
    confirm = next(button for button in warning.buttons if button.value.startswith("admin:remove:confirm:"))
    accepted = await engine.handle(inbound(8, confirm.value, callback=True))

    assert accepted.deferred and accepted.command_enqueued
    command = publisher.commands[-1]
    assert command.operation is Operation.REMOVE_PARTICIPANT
    assert command.participant_key == registrations[1].participant_key
    assert command.reply_context.text_success == (f"🗑 Игрок «Фамилия1 Имя1» удалён из игры «{event.name}».")


@pytest.mark.asyncio
async def test_event_leader_confirms_exact_waiting_player_count_before_sifting(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    engine, _, publisher, tables = await engine_setup(disk_store, event, admin=True)
    await tables.import_missing(
        [
            Registration(
                event_id=event.event_id,
                participant_key="a" * 43,
                display_name="Ожидающий Один",
                wish_play="Без пожеланий",
            ),
            Registration(
                event_id=event.event_id,
                participant_key="b" * 43,
                display_name="Ожидающий Два",
                wish_play="Без пожеланий",
            ),
            Registration(
                event_id=event.event_id,
                participant_key="c" * 43,
                display_name="Подтвердивший Игрок",
                wish_play="Без пожеланий",
                attendance_status=AttendanceStatus.CONFIRMED,
            ),
        ]
    )

    warning = await engine.handle(inbound(1, f"ag:remove-waiting:{event.event_id}", callback=True))

    assert "Не подтвердили участие: 2." in warning.text
    assert "получат уведомления" in warning.text
    confirm = next(button for button in warning.buttons if button.value.startswith("admin:remove-waiting:confirm:"))
    assert confirm.label == "Да, пересортировать"
    accepted = await engine.handle(inbound(2, confirm.value, callback=True))

    assert accepted.deferred and accepted.command_enqueued
    command = publisher.commands[-1]
    assert command.operation is Operation.REMOVE_WAITING_PARTICIPANTS
    assert command.participant_key is None
    assert command.reply_context.text_success == ""


@pytest.mark.asyncio
@pytest.mark.parametrize("profile", ["@Target_Player", "https://t.me/Target_Player"])
async def test_admin_grants_telegram_gamemaster_by_profile_and_notifies_them(
    profile: str, disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, admin=True)
    await users.save(
        TelegramUser(
            tg_id=2,
            telegram_handle="@target_player",
            full_name="Петрова Анна",
            vk_url="https://vk.com/anna",
            crossplay=True,
            larp_experience=False,
            needs_pass=False,
        )
    )

    menu = await engine.handle(inbound(1, ADMIN))
    assert MANAGE_MODERATORS in [button.value for button in menu.buttons]
    moderator_menu = await engine.handle(inbound(2, MANAGE_MODERATORS, callback=True))
    assert GRANT_GAMEMASTER in [button.value for button in moderator_menu.buttons]
    choose_platform = await engine.handle(inbound(3, GRANT_GAMEMASTER, callback=True))
    assert [button.label for button in choose_platform.buttons[:2]] == ["Telegram", "VK"]
    prompt = await engine.handle(inbound(4, "admin:gamemaster:telegram", callback=True))
    assert "ссылку на профиль Telegram или @username" in prompt.text
    granted = await engine.handle(inbound(5, profile))

    target = await users.get(Platform.TELEGRAM, 2)
    assert target is not None and target.is_gamemaster
    assert target.gamemaster_grant_operation_id == "gamemaster-grant:telegram:2:5"
    assert isinstance(engine.transport, MemoryDeferredTransport)
    assert engine.transport.sent == [
        (
            Platform.TELEGRAM,
            2,
            "gamemaster-grant:telegram:2:5",
            GAMEMASTER_NOTIFICATION,
            (Button(label=ADMIN_MENU, value=ADMIN_MENU),),
        ),
    ]
    assert target.last_bot_buttons == [Button(label=ADMIN_MENU, value=ADMIN_MENU)]
    assert granted.text == "✅ Пользователь назначен модератором в Telegram."

    target_main = await engine.handle(inbound(6, "/start", user_id=2, telegram_username="target_player"))
    assert ADMIN in [button.value for button in target_main.buttons]


@pytest.mark.asyncio
@pytest.mark.parametrize("profile", ["@legacy_player", "[id42|@legacy_player]"])
async def test_admin_grants_legacy_telegram_user_via_registration_history_and_backfills_handle(
    profile: str, disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, registrations = await engine_setup(disk_store, event, admin=True)
    await users.save(
        TelegramUser(
            tg_id=2,
            full_name="Петрова Анна",
            vk_url="https://vk.com/old_alias",
            crossplay=True,
            larp_experience=False,
            needs_pass=False,
        )
    )
    await registrations.enlist(
        event.event_id,
        operation_id="legacy-enlist",
        participant_key=engine.registrations.key(Platform.TELEGRAM, 2, event.event_id),
        display_name="Петрова Анна",
        telegram_profile="https://t.me/legacy_player",
        wish_play="Anyone",
    )

    await engine.handle(inbound(1, GRANT_GAMEMASTER))
    await engine.handle(inbound(2, "admin:gamemaster:telegram", callback=True))
    granted = await engine.handle(inbound(3, profile))

    target = await users.get(Platform.TELEGRAM, 2)
    assert target is not None and target.is_gamemaster
    assert target.telegram_handle == "@legacy_player"
    assert granted.text == "✅ Пользователь назначен модератором в Telegram."


@pytest.mark.asyncio
async def test_admin_does_not_grant_when_registration_history_matches_multiple_telegram_users(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, registrations = await engine_setup(disk_store, event, admin=True)
    for target_id in (2, 3):
        await users.save(
            TelegramUser(
                tg_id=target_id,
                full_name=f"Игрок {target_id}",
                vk_url=f"https://vk.com/id{target_id}",
                crossplay=True,
                larp_experience=False,
                needs_pass=False,
            )
        )
        await registrations.enlist(
            event.event_id,
            operation_id=f"legacy-enlist-{target_id}",
            participant_key=engine.registrations.key(Platform.TELEGRAM, target_id, event.event_id),
            display_name=f"Игрок {target_id}",
            telegram_profile="https://t.me/reused_handle",
            wish_play="Anyone",
        )

    await engine.handle(inbound(1, GRANT_GAMEMASTER))
    await engine.handle(inbound(2, "admin:gamemaster:telegram", callback=True))
    rejected = await engine.handle(inbound(3, "@reused_handle"))

    assert rejected.text.startswith("Пользователь Telegram не найден.")
    for target_id in (2, 3):
        target = await users.get(Platform.TELEGRAM, target_id)
        assert target is not None and not target.is_gamemaster and target.telegram_handle is None


@pytest.mark.asyncio
async def test_admin_grants_vk_gamemaster_from_vanity_link(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, users, _, _ = await engine_setup(
        disk_store,
        event,
        admin=True,
        vk_profile_ids={"https://vk.com/game_master": 22},
    )
    await users.save(
        VkUser(
            vk_id=22,
            full_name="Петров Пётр",
            crossplay=False,
            larp_experience=True,
            needs_pass=False,
        )
    )

    await engine.handle(inbound(1, GRANT_GAMEMASTER))
    await engine.handle(inbound(2, "admin:gamemaster:vk", callback=True))
    granted = await engine.handle(inbound(3, "vk.ru/game_master"))

    target = await users.get(Platform.VK, 22)
    assert target is not None and target.is_gamemaster
    assert isinstance(engine.transport, MemoryDeferredTransport)
    assert engine.transport.sent == [
        (
            Platform.VK,
            22,
            "gamemaster-grant:vk:22:3",
            GAMEMASTER_NOTIFICATION,
            (Button(label=ADMIN_MENU, value=ADMIN_MENU),),
        ),
    ]
    assert target.last_bot_buttons == [Button(label=ADMIN_MENU, value=ADMIN_MENU)]
    assert granted.text == "✅ Пользователь назначен модератором в VK."


@pytest.mark.asyncio
async def test_admin_rejects_incomplete_gamemaster_profile(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, admin=True)
    await users.save(TelegramUser(tg_id=2, telegram_handle="@target_player"))

    await engine.handle(inbound(1, GRANT_GAMEMASTER))
    await engine.handle(inbound(2, "admin:gamemaster:telegram", callback=True))
    rejected = await engine.handle(inbound(3, "@target_player"))

    target = await users.get(Platform.TELEGRAM, 2)
    assert target is not None and not target.is_gamemaster
    assert "заполнен не полностью" in rejected.text
    assert isinstance(engine.transport, MemoryDeferredTransport)
    assert engine.transport.sent == []


@pytest.mark.asyncio
async def test_gamemaster_notification_retries_after_role_was_persisted(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, admin=True)
    engine.transport = FailOnceTransport()
    await users.save(
        TelegramUser(
            tg_id=2,
            telegram_handle="@target_player",
            full_name="Петрова Анна",
            vk_url="https://vk.com/anna",
            crossplay=True,
            larp_experience=False,
            needs_pass=False,
        )
    )
    await engine.handle(inbound(1, GRANT_GAMEMASTER))
    await engine.handle(inbound(2, "admin:gamemaster:telegram", callback=True))
    grant_message = inbound(3, "@target_player")

    with pytest.raises(RuntimeError, match="temporary transport failure"):
        await engine.handle(grant_message)
    persisted = await users.get(Platform.TELEGRAM, 2)
    assert persisted is not None and persisted.is_gamemaster
    assert persisted.last_delivery_operation_id is None

    response = await engine.handle(grant_message)

    assert response.text == "✅ Пользователь назначен модератором в Telegram."
    assert isinstance(engine.transport, FailOnceTransport)
    assert len(engine.transport.sent) == 1
    assert engine.transport.sent[0][3] == GAMEMASTER_NOTIFICATION
    assert engine.transport.sent[0][4] == (Button(label=ADMIN_MENU, value=ADMIN_MENU),)


@pytest.mark.asyncio
async def test_gamemaster_cannot_grant_gamemaster_role(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, _, _ = await engine_setup(disk_store, event, gamemaster=True)

    response = await engine.handle(inbound(1, GRANT_GAMEMASTER))

    assert response.text == "Недостаточно прав."


@pytest.mark.asyncio
async def test_admin_lists_and_removes_moderators_with_pagination_and_notification(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, admin=True)
    for target_id in range(2, 14):
        await users.save(
            TelegramUser(
                tg_id=target_id,
                telegram_handle=f"@moderator_{target_id}",
                full_name=f"Модератор {target_id:02d}",
                vk_url=f"https://vk.com/id{target_id}",
                crossplay=True,
                larp_experience=False,
                needs_pass=False,
                is_gamemaster=True,
            )
        )

    menu = await engine.handle(inbound(1, ADMIN))
    assert MANAGE_MODERATORS in {button.value for button in menu.buttons}
    role_menu = await engine.handle(inbound(2, MANAGE_MODERATORS, callback=True))
    assert {GRANT_GAMEMASTER, REVOKE_GAMEMASTER, LIST_GAMEMASTERS} <= {button.value for button in role_menu.buttons}

    first_page = await engine.handle(inbound(3, LIST_GAMEMASTERS, callback=True))
    profile_buttons = [button for button in first_page.buttons if button.value.startswith("mod:view:")]
    assert len(profile_buttons) == PAGE_SIZE
    assert any(button.value == f"mod:list-page:{PAGE_SIZE}" for button in first_page.buttons)
    profile = await engine.handle(inbound(4, profile_buttons[0].value, callback=True))
    assert "Модератор" in profile.text
    assert "https://t.me/moderator_" in profile.text
    assert "https://vk.com/id" in profile.text

    returned = await engine.handle(inbound(5, profile.buttons[0].value, callback=True))
    await engine.handle(
        inbound(
            6, next(button.value for button in returned.buttons if button.value == MANAGE_MODERATORS), callback=True
        )
    )
    removal = await engine.handle(inbound(7, REVOKE_GAMEMASTER, callback=True))
    target_button = next(button for button in removal.buttons if button.value == "mod:remove-pick:t:2")
    warning = await engine.handle(inbound(8, target_button.value, callback=True))
    assert "Снять полномочия модератора" in warning.text
    removed = await engine.handle(inbound(9, "mod:remove-confirm", callback=True))

    target = await users.get(Platform.TELEGRAM, 2)
    assert target is not None and not target.is_gamemaster
    assert removed.text == "✅ Полномочия модератора сняты."
    assert isinstance(engine.transport, MemoryDeferredTransport)
    assert engine.transport.sent[-1][1:] == (
        2,
        "gamemaster-revoke:telegram:2:9",
        GAMEMASTER_REVOKED_NOTIFICATION,
        (),
    )


@pytest.mark.asyncio
async def test_created_game_confirmation_includes_all_table_links(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, admin=True)
    creator = await users.get(Platform.TELEGRAM, 1)
    assert isinstance(creator, TelegramUser)
    creator.needs_pass = True
    creator.pass_details = PassDetails(
        surname_cyrillic="Иванов",
        name_cyrillic="Иван",
        patronym_cyrillic="-",
        foreigner=False,
        mobile_phone="+7 999 123-45-67",
        email="creator@example.com",
    )
    await users.save(creator)

    prompt = await engine.handle(inbound(1, "➕ Создать игру"))
    assert prompt.text == "Введите название игры:"
    amount_prompt = await engine.handle(inbound(2, "Бал в зимнем дворце"))
    assert "DD.MM.YY" in amount_prompt.text
    assert amount_prompt.buttons == [Button(label="Следующая суббота", value="admin:date:next-saturday")]
    invalid_date = await engine.handle(inbound(3, "31.02.26"))
    assert "Некорректная дата" in invalid_date.text
    assert invalid_date.buttons == amount_prompt.buttons
    amount_prompt = await engine.handle(inbound(4, "14.11.26"))
    assert amount_prompt.text == "Введите количество игроков (от 1 до 10000):"
    invalid = await engine.handle(inbound(5, "много"))
    assert invalid.text == "Введите целое количество игроков от 1 до 10000:"
    question = await engine.handle(inbound(6, "24"))
    assert question.text == "Собираетесь ли вы сами мастерить эту игру, или добавите на эту роль кого-то ещё?"
    assert question.buttons == [
        Button(label="Да, буду мастерить самостоятельно", value="create:master-self"),
        Button(label="Нет, выбрать мастера", value="create:master-pick"),
    ]
    assert [candidate.name for candidate in await engine.events.list_page(limit=10)] == [event.name]
    response = await engine.handle(inbound(7, "create:master-self", callback=True))

    created_events = await engine.events.list_page(limit=10)
    created = next(candidate for candidate in created_events if candidate.name == "Бал в зимнем дворце")
    assert created.master_table_resource_path.endswith("/master_table_Бал в зимнем дворце.xlsx")
    assert created.game_date == date(2026, 11, 14)
    assert created.player_amount == 24
    assert created.public_table_resource_path is not None
    assert created.public_table_resource_path.endswith("/public_table_Бал в зимнем дворце.xlsx")
    assert created.pass_table_resource_path == f"disk:/larp-bot/passes/{created.event_id}.xlsx"
    assert created.pass_table_public_url is not None
    for path in (created.master_table_resource_path, created.public_table_resource_path):
        workbook = load_workbook(BytesIO(disk_store.files[path]))
        try:
            assert workbook.active.cell(26, 1).value == "ЗАПАС"
        finally:
            workbook.close()
    pass_workbook = load_workbook(BytesIO(disk_store.files[created.pass_table_resource_path]))
    try:
        assert pass_workbook.active.cell(2, 1).value == "Иванов"
    finally:
        pass_workbook.close()
    assert response.text.count("https://disk.example/public/") == 3
    assert MASTER_TABLE_DISCLAIMER in response.text
    assert "Публичная таблица (без контактов Telegram и VK)" in response.text
    assert "Административная таблица" in response.text
    assert "Таблица пропусков" in response.text
    assert response.buttons == [Button(label="⬅️ Назад", value=f"ag:manage:{created.event_id}")]
    assert await engine.events.list_leaders(created.event_id) == [
        BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1)
    ]

    assert "Дата игры: 14.11.26" in response.text
    management = await engine.handle(inbound(8, response.buttons[0].value, callback=True))
    assert "Управление игрой «Бал в зимнем дворце» (14.11.26)" in management.text
    assert EVENT_TABLES in [button.label for button in management.buttons]
    tables = await engine.handle(inbound(9, f"ag:tables:{created.event_id}", callback=True))
    assert tables.text.count("https://disk.example/public/") == 3
    assert MASTER_TABLE_DISCLAIMER in tables.text
    assert "Публичная таблица (без контактов Telegram и VK)" in tables.text
    assert "Административная таблица" in tables.text
    assert "Таблица пропусков" in tables.text


@pytest.mark.asyncio
async def test_admin_can_choose_next_saturday_as_game_date(
    monkeypatch: pytest.MonkeyPatch,
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    next_saturday = date(2026, 9, 19)
    monkeypatch.setattr("larp_bot.application.conversation.next_saturday_game_date", lambda: next_saturday)
    engine, users, _, _ = await engine_setup(disk_store, event, admin=True)

    await engine.handle(inbound(1, "➕ Создать игру"))
    date_prompt = await engine.handle(inbound(2, "Игра в субботу"))
    amount_prompt = await engine.handle(inbound(3, date_prompt.buttons[0].value, callback=True))

    creator = await users.get(Platform.TELEGRAM, 1)
    assert creator is not None
    assert amount_prompt.text == "Введите количество игроков (от 1 до 10000):"
    assert creator.dialog_context["game_date"] == next_saturday.isoformat()


@pytest.mark.asyncio
async def test_gamemaster_created_game_seeds_only_its_creator(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, _, _ = await engine_setup(
        disk_store,
        event,
        gamemaster=True,
        tg_admin_ids={20},
        vk_admin_ids={30},
    )

    await engine.handle(inbound(1, "➕ Создать игру"))
    await engine.handle(inbound(2, "Новая игра"))
    await engine.handle(inbound(3, "15.11.26"))
    response = await engine.handle(inbound(4, "20"))

    assert response.text.startswith("✅ Игра создана.")
    assert response.text.count("https://disk.example/public/") == 3
    created = next(candidate for candidate in await engine.events.list_page(limit=10) if candidate.name == "Новая игра")
    leaders = await engine.events.list_leaders(created.event_id)
    assert {(leader.platform, leader.platform_user_id) for leader in leaders} == {
        (Platform.TELEGRAM, 1),
    }


@pytest.mark.asyncio
async def test_admin_created_game_can_pick_another_moderator_or_admin_as_its_first_master(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, admin=True, tg_admin_ids={20})
    for target_id in (*range(2, 8), 20):
        await users.save(
            TelegramUser(
                tg_id=target_id,
                telegram_handle=f"@privileged_{target_id}",
                full_name=f"Мастер {target_id:02d}",
                vk_url=f"https://vk.com/id{target_id}",
                crossplay=True,
                larp_experience=False,
                needs_pass=False,
                is_gamemaster=target_id != 20,
            )
        )

    await engine.handle(inbound(1, "➕ Создать игру"))
    await engine.handle(inbound(2, "Новая игра"))
    await engine.handle(inbound(3, "15.11.26"))
    await engine.handle(inbound(4, "20"))
    first_page = await engine.handle(inbound(5, "create:master-pick", callback=True))
    second_page = await engine.handle(inbound(6, "create:master-page:5", callback=True))

    assert first_page.text == (
        "Выберите модератора или администратора, который станет первым мастером игры «Новая игра» (15.11.26):"
    )
    assert [button.value for button in first_page.buttons] == [
        *(f"create:master:t:{target_id}" for target_id in range(2, 7)),
        "create:master-page:5",
        "create:master-choice",
    ]
    assert [button.label for button in second_page.buttons] == [
        "Мастер 07 · TG",
        "Мастер 20 · TG",
        PAGE_BACK,
        "Отмена",
    ]
    assert [candidate.name for candidate in await engine.events.list_page(limit=10)] == [event.name]

    response = await engine.handle(inbound(7, "create:master:t:20", callback=True))

    created = next(candidate for candidate in await engine.events.list_page(limit=10) if candidate.name == "Новая игра")
    assert await engine.events.list_leaders(created.event_id) == [
        BotIdentity(platform=Platform.TELEGRAM, platform_user_id=20)
    ]
    assert response.text.startswith("✅ Игра создана.")
    assert "Мастер: Мастер 20 · TG" in response.text
    assert response.text.count("https://disk.example/public/") == 3
    assert response.buttons == [Button(label=BACK, value=f"ag:manage:{created.event_id}")]
    assert isinstance(engine.transport, MemoryDeferredTransport)
    assert engine.transport.sent == [
        (
            Platform.TELEGRAM,
            20,
            f"leader-grant:{created.event_id}:telegram:20:7",
            "Вы добавлены мастером игры «Новая игра» (15.11.26).",
            (Button(label="Управление игрой", value=f"ag:manage:{created.event_id}"),),
        )
    ]


@pytest.mark.asyncio
async def test_admin_created_game_master_notification_retry_does_not_create_a_second_game(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, admin=True)
    engine.transport = FailOnceTransport()
    await users.save(
        TelegramUser(
            tg_id=2,
            telegram_handle="@moderator_2",
            full_name="Петрова Анна",
            vk_url="https://vk.com/anna",
            crossplay=True,
            larp_experience=False,
            needs_pass=False,
            is_gamemaster=True,
        )
    )
    await engine.handle(inbound(1, "➕ Создать игру"))
    await engine.handle(inbound(2, "Новая игра"))
    await engine.handle(inbound(3, "15.11.26"))
    await engine.handle(inbound(4, "20"))
    await engine.handle(inbound(5, "create:master-pick", callback=True))
    pick = inbound(6, "create:master:t:2", callback=True)

    with pytest.raises(RuntimeError, match="temporary transport failure"):
        await engine.handle(pick)
    response = await engine.handle(pick)

    created = [candidate for candidate in await engine.events.list_page(limit=10) if candidate.name == "Новая игра"]
    assert len(created) == 1
    assert await engine.events.list_leaders(created[0].event_id) == [
        BotIdentity(platform=Platform.TELEGRAM, platform_user_id=2)
    ]
    assert response.text.startswith("✅ Игра создана.")
    assert isinstance(engine.transport, FailOnceTransport)
    assert [sent[2] for sent in engine.transport.sent] == [f"leader-grant:{created[0].event_id}:telegram:2:6"]


@pytest.mark.asyncio
async def test_admin_creation_master_picker_returns_to_question_and_rechecks_admin_role(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, _, _, _ = await engine_setup(disk_store, event, admin=True)
    await engine.handle(inbound(1, "➕ Создать игру"))
    await engine.handle(inbound(2, "Новая игра"))
    await engine.handle(inbound(3, "15.11.26"))
    await engine.handle(inbound(4, "20"))

    empty = await engine.handle(inbound(5, "create:master-pick", callback=True))
    question = await engine.handle(inbound(6, "create:master-choice", callback=True))

    assert empty.text == "Нет других модераторов или администраторов, которых можно выбрать мастером."
    assert empty.buttons == [Button(label="Отмена", value="create:master-choice")]
    assert question.text == CREATION_MASTER_QUESTION

    assert isinstance(engine.admins, StaticAdminProvider)
    engine.admins.tg_ids.clear()
    engine.admins.tg_gamemaster_ids.add(1)
    denied = await engine.handle(inbound(7, "create:master-self", callback=True))

    assert denied.text == "Недостаточно прав."
    assert [candidate.name for candidate in await engine.events.list_page(limit=10)] == [event.name]


@pytest.mark.asyncio
async def test_admin_controls_existing_game_without_becoming_its_master(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, _, _, _ = await engine_setup(
        disk_store,
        event,
        admin=True,
        tg_admin_ids={20},
        vk_admin_ids={30},
    )

    games = await engine.handle(inbound(1, MANAGE_GAMES))
    selected = await engine.handle(inbound(2, f"ag:manage:{event.event_id}", callback=True))

    assert event.name in next(button.label for button in games.buttons if button.value.startswith("ag:manage:"))
    assert f"Управление игрой «{event.name}»" in selected.text
    assert f"Количество игроков: {event.player_amount}" in selected.text
    assert CHANGE_STATUS in {button.label for button in selected.buttons}
    assert await engine.events.list_leaders(event.event_id) == []


@pytest.mark.asyncio
async def test_game_master_can_change_player_amount(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, publisher, _ = await engine_setup(disk_store, event, gamemaster=True)
    await engine.events.add_leader(
        event.event_id,
        BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1),
    )

    management = await engine.handle(inbound(1, f"ag:manage:{event.event_id}", callback=True))
    amount_button = next(button for button in management.buttons if button.label == CHANGE_PLAYER_AMOUNT)
    prompt = await engine.handle(inbound(2, amount_button.value, callback=True))
    invalid = await engine.handle(inbound(3, "0"))
    accepted = await engine.handle(inbound(4, "24"))

    assert prompt.text == (
        f"Текущее количество игроков: {event.player_amount}.\n\nВведите новое количество игроков (от 1 до 10000):"
    )
    assert invalid.text == "Введите целое количество игроков от 1 до 10000:"
    assert accepted.text == "⏳ Изменение количества игроков принято в обработку."
    assert accepted.deferred and accepted.command_enqueued
    queued = publisher.commands[-1]
    assert queued.operation is Operation.UPDATE_PLAYER_AMOUNT
    assert queued.participant_key is None
    assert queued.payload == PlayerAmountPayload(player_amount=24)
    assert queued.reply_context.text_success == f"Количество игроков игры «{event.name}» изменено: 24."


@pytest.mark.asyncio
async def test_admin_can_join_and_leave_game_masters_without_losing_control_or_pass_sync(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, admin=True)
    admin = await users.get(Platform.TELEGRAM, 1)
    assert isinstance(admin, TelegramUser)
    admin.needs_pass = True
    admin.pass_details = PassDetails(
        surname_cyrillic="Иванов",
        name_cyrillic="Иван",
        patronym_cyrillic="-",
        foreigner=False,
        mobile_phone="+7 999 123-45-67",
        email="admin@example.com",
    )
    await users.save(admin)

    await engine.administration.create_pass_table(event.event_id)
    stored_event = await engine.events.get(event.event_id)
    assert stored_event is not None and stored_event.pass_table_resource_path is not None
    pass_path = stored_event.pass_table_resource_path
    initial = load_workbook(BytesIO(disk_store.files[pass_path]))
    try:
        assert initial.active.max_row == 1
    finally:
        initial.close()

    management = await engine.handle(inbound(1, f"ag:manage:{event.event_id}", callback=True))
    assert CHANGE_STATUS in {button.label for button in management.buttons}
    masters = await engine.handle(inbound(2, f"ag:masters:{event.event_id}", callback=True))
    self_add = next(button for button in masters.buttons if button.label == ADD_SELF_EVENT_LEADER)
    added = await engine.handle(inbound(3, self_add.value, callback=True))

    assert added.text == f"✅ Администратор добавлен мастером игры «{event.name}»."
    identity = BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1)
    assert identity in await engine.events.list_leaders(event.event_id)
    joined = load_workbook(BytesIO(disk_store.files[pass_path]))
    try:
        assert joined.active.max_row == 2
        assert joined.active["I2"].value == "admin@example.com"
    finally:
        joined.close()

    await engine.handle(inbound(4, f"ag:masters:{event.event_id}", callback=True))
    removal = await engine.handle(inbound(5, f"master:remove:{event.event_id}", callback=True))
    remove_self = next(button for button in removal.buttons if button.value == "master:remove-pick:t:1")
    await engine.handle(inbound(6, remove_self.value, callback=True))
    removed = await engine.handle(inbound(7, "master:remove-confirm", callback=True))

    assert removed.text == f"✅ Мастер удалён из игры «{event.name}»."
    assert identity not in await engine.events.list_leaders(event.event_id)
    left = load_workbook(BytesIO(disk_store.files[pass_path]))
    try:
        assert left.active.max_row == 1
    finally:
        left.close()

    still_admin = await engine.handle(inbound(8, f"ag:manage:{event.event_id}", callback=True))
    assert CHANGE_STATUS in {button.label for button in still_admin.buttons}


@pytest.mark.asyncio
async def test_event_leader_can_add_another_gamemaster_who_can_then_notify_players(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, publisher, _ = await engine_setup(disk_store, event, gamemaster=True)
    await engine.events.add_leader(
        event.event_id,
        BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1),
    )
    await users.save(
        TelegramUser(
            tg_id=2,
            telegram_handle="@target_player",
            full_name="Петрова Анна",
            vk_url="https://vk.com/anna",
            crossplay=True,
            larp_experience=False,
            needs_pass=False,
            is_gamemaster=True,
        )
    )

    await engine.handle(inbound(1, MANAGE_GAMES))
    await engine.handle(inbound(2, f"ag:manage:{event.event_id}", callback=True))
    platforms = await engine.handle(inbound(3, f"ag:leader:{event.event_id}", callback=True))
    assert [button.label for button in platforms.buttons[:2]] == ["Telegram", "VK"]
    await engine.handle(inbound(4, "admin:event-leader:telegram", callback=True))
    added = await engine.handle(inbound(5, "@target_player"))

    assert added.text == f"✅ Модератор добавлен мастером игры «{event.name}»."
    assert BotIdentity(platform=Platform.TELEGRAM, platform_user_id=2) in await engine.events.list_leaders(
        event.event_id
    )

    await engine.handle(inbound(1, MANAGE_GAMES, user_id=2))
    await engine.handle(inbound(2, f"ag:manage:{event.event_id}", callback=True, user_id=2))
    await engine.handle(inbound(3, f"ag:notify:{event.event_id}", callback=True, user_id=2))
    preview = await engine.handle(inbound(4, "Сбор в 18:30", user_id=2))
    queued = await engine.handle(inbound(5, preview.buttons[0].value, callback=True, user_id=2))

    assert queued.deferred
    assert publisher.commands[-1].operation is Operation.SEND_CONFIRMED_NOTIFICATION
    assert publisher.commands[-1].platform_user_id == 2


@pytest.mark.asyncio
async def test_master_management_adds_lists_and_removes_moderator_with_notifications(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, gamemaster=True)
    await engine.events.add_leader(
        event.event_id,
        BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1),
    )
    for target_id in range(2, 14):
        await users.save(
            TelegramUser(
                tg_id=target_id,
                telegram_handle=f"@moderator_{target_id}",
                full_name=f"Модератор {target_id:02d}",
                vk_url=f"https://vk.com/id{target_id}",
                crossplay=True,
                larp_experience=False,
                needs_pass=False,
                is_gamemaster=True,
            )
        )

    management = await engine.handle(inbound(1, f"ag:manage:{event.event_id}", callback=True))
    masters_button = next(button for button in management.buttons if button.label == MANAGE_MASTERS)
    masters = await engine.handle(inbound(2, masters_button.value, callback=True))
    assert {ADD_EVENT_LEADER, REMOVE_EVENT_LEADER, LIST_EVENT_LEADERS} <= {button.label for button in masters.buttons}

    candidates = await engine.handle(inbound(3, f"master:add:{event.event_id}", callback=True))
    candidate_buttons = [button for button in candidates.buttons if button.value.startswith("master:add-pick:")]
    assert len(candidate_buttons) == PAGE_SIZE
    assert any(button.value == f"master:add-page:{PAGE_SIZE}" for button in candidates.buttons)
    assert any(button.value == "master:add-direct" for button in candidates.buttons)
    added = await engine.handle(inbound(4, "master:add-pick:t:2", callback=True))

    assert added.text == f"✅ Модератор добавлен мастером игры «{event.name}»."
    assert isinstance(engine.transport, MemoryDeferredTransport)
    assert engine.transport.sent[-1] == (
        Platform.TELEGRAM,
        2,
        f"leader-grant:{event.event_id}:telegram:2:4",
        f"Вы добавлены мастером игры «{event.name}».",
        (Button(label="Управление игрой", value=f"ag:manage:{event.event_id}"),),
    )

    await engine.handle(inbound(5, f"ag:masters:{event.event_id}", callback=True))
    listing = await engine.handle(inbound(6, f"master:list:{event.event_id}", callback=True))
    target_button = next(button for button in listing.buttons if button.value == "master:view:t:2")
    profile = await engine.handle(inbound(7, target_button.value, callback=True))
    assert profile.text.startswith("Мастер\n")
    assert "https://t.me/moderator_2" in profile.text
    assert "https://vk.com/id2" in profile.text

    returned = await engine.handle(inbound(8, profile.buttons[0].value, callback=True))
    await engine.handle(
        inbound(
            9,
            next(button.value for button in returned.buttons if button.value == f"ag:masters:{event.event_id}"),
            callback=True,
        )
    )
    removal = await engine.handle(inbound(10, f"master:remove:{event.event_id}", callback=True))
    remove_button = next(button for button in removal.buttons if button.value == "master:remove-pick:t:2")
    warning = await engine.handle(inbound(11, remove_button.value, callback=True))
    assert "Удалить мастера" in warning.text
    removed = await engine.handle(inbound(12, "master:remove-confirm", callback=True))

    assert removed.text == f"✅ Мастер удалён из игры «{event.name}»."
    assert BotIdentity(platform=Platform.TELEGRAM, platform_user_id=2) not in await engine.events.list_leaders(
        event.event_id
    )
    assert engine.transport.sent[-1] == (
        Platform.TELEGRAM,
        2,
        f"leader-revoke:{event.event_id}:telegram:2:12",
        f"Вы больше не являетесь мастером игры «{event.name}».",
        (),
    )


@pytest.mark.asyncio
async def test_event_leader_cannot_add_a_non_gamemaster_as_leader(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, users, _, _ = await engine_setup(disk_store, event, gamemaster=True)
    await engine.events.add_leader(
        event.event_id,
        BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1),
    )
    await users.save(
        TelegramUser(
            tg_id=2,
            telegram_handle="@ordinary_user",
            full_name="Петрова Анна",
            vk_url="https://vk.com/anna",
            crossplay=True,
            larp_experience=False,
            needs_pass=False,
        )
    )

    await engine.handle(inbound(1, f"ag:leader:{event.event_id}", callback=True))
    await engine.handle(inbound(2, "admin:event-leader:telegram", callback=True))
    denied = await engine.handle(inbound(3, "@ordinary_user"))

    assert denied.text == "Добавить мастером можно только модератора или администратора."
    assert BotIdentity(platform=Platform.TELEGRAM, platform_user_id=2) not in await engine.events.list_leaders(
        event.event_id
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("archive_state", ["ADMIN_ARCHIVE_NAME", "ADMIN_DELETE_NAME"])
async def test_nonleader_gamemaster_cannot_archive_through_current_or_stale_actions(
    archive_state: str, disk_store: MemoryDiskStore, event: Event
) -> None:
    engine, users, publisher, _ = await engine_setup(disk_store, event, gamemaster=True)

    listing = await engine.handle(inbound(1, MANAGE_GAMES))
    assert any(button.value == f"ag:manage:{event.event_id}" for button in listing.buttons)
    current = await engine.handle(inbound(2, f"ag:archive:{event.event_id}", callback=True))
    assert current.text == "Только мастера этой игры могут изменять её данные."
    for update, value in enumerate(
        (
            f"select:admin-archive:{event.event_id}",
            f"select:admin-delete:{event.event_id}",
        ),
        start=3,
    ):
        response = await engine.handle(inbound(update, value, callback=True))
        assert response.text == "Только мастера этой игры могут изменять её данные."

    page = await engine.handle(
        inbound(
            5,
            f"page:admin-archive:{int(event.created_at.timestamp() * 1_000_000)}:{event.event_id}",
            callback=True,
        )
    )
    assert "мастера" not in page.text

    user = await users.get(Platform.TELEGRAM, 1)
    assert user is not None
    user.dialog_state = archive_state
    user.dialog_context = {"event_id": event.event_id, "event_name": event.name}
    await users.save(user)
    response = await engine.handle(inbound(10, event.name))

    assert response.text == "Только мастера этой игры могут изменять её данные."
    saved_user = await users.get(Platform.TELEGRAM, 1)
    assert saved_user is not None and saved_user.dialog_state == "IDLE"
    assert publisher.commands == []


@pytest.mark.asyncio
async def test_admin_game_buttons_are_newest_first_and_paginated(disk_store: MemoryDiskStore) -> None:
    base = datetime(2026, 1, 1, tzinfo=UTC)
    events = [
        Event(
            event_id=f"event-{index:04d}",
            name=f"Game {index}",
            disk_resource_path=f"disk:/larp-bot/events/event-{index:04d}-game.xlsx",
            public_registration_url=f"https://disk.example/{index}",
            created_at=base + timedelta(seconds=index),
            updated_at=base + timedelta(seconds=index),
        )
        for index in range(21)
    ]
    engine, _, _, _ = await engine_setup(disk_store, *events, admin=True)

    first = await engine.handle(inbound(1, MANAGE_GAMES))
    game_buttons = [button for button in first.buttons if button.value.startswith("ag:manage:")]
    next_button = next(button for button in first.buttons if button.label == PAGE_FORWARD)
    assert len(game_buttons) == PAGE_SIZE
    assert [button.label for button in game_buttons] == [f"Game {index}" for index in range(20, 15, -1)]
    assert PAGE_BACK not in {button.label for button in first.buttons}

    second = await engine.handle(inbound(2, next_button.value, callback=True))
    assert [button.label for button in second.buttons if button.value.startswith("ag:manage:")] == [
        f"Game {index}" for index in range(15, 10, -1)
    ]
    previous_button = next(button for button in second.buttons if button.label == PAGE_BACK)
    returned = await engine.handle(inbound(3, previous_button.value, callback=True))
    assert [button.label for button in returned.buttons if button.value.startswith("ag:manage:")] == [
        f"Game {index}" for index in range(20, 15, -1)
    ]
    second_next = next(button for button in second.buttons if button.label == PAGE_FORWARD)
    third = await engine.handle(inbound(4, second_next.value, callback=True))
    assert [button.label for button in third.buttons if button.value.startswith("ag:manage:")] == [
        f"Game {index}" for index in range(10, 5, -1)
    ]
    second_again = await engine.handle(
        inbound(5, next(button.value for button in third.buttons if button.label == PAGE_BACK), callback=True)
    )
    assert [button.label for button in second_again.buttons if button.value.startswith("ag:manage:")] == [
        f"Game {index}" for index in range(15, 10, -1)
    ]
    fourth = await engine.handle(
        inbound(6, next(button.value for button in third.buttons if button.label == PAGE_FORWARD), callback=True)
    )
    last = await engine.handle(
        inbound(7, next(button.value for button in fourth.buttons if button.label == PAGE_FORWARD), callback=True)
    )
    assert [button.label for button in last.buttons if button.value.startswith("ag:manage:")] == ["Game 0"]
    assert PAGE_BACK in {button.label for button in last.buttons}
    assert PAGE_FORWARD not in {button.label for button in last.buttons}
    assert last.buttons[-1] == Button(label="Отмена", value=ADMIN_MENU)


@pytest.mark.asyncio
async def test_admin_can_unarchive_game_from_archive_menu(disk_store: MemoryDiskStore, event: Event) -> None:
    now = datetime(2026, 1, 1, tzinfo=UTC)
    closed = Event(
        event_id="event-closed",
        name="Closed but managed",
        disk_resource_path="disk:/larp-bot/events/event-closed-game.xlsx",
        public_registration_url="https://disk.example/closed",
        status=EventStatus.CLOSED,
        created_at=now,
        updated_at=now,
    )
    archived = [
        Event.model_validate(
            event.model_dump()
            | {
                "event_id": f"event-archived-{index}",
                "name": f"Archived {index}",
                "status": EventStatus.CLOSED,
                "created_at": now + timedelta(seconds=index + 1),
                "archived_at": now,
            }
        )
        for index in range(PAGE_SIZE + 1)
    ]
    engine, _, publisher, _ = await engine_setup(disk_store, closed, *archived, admin=True)

    games = await engine.handle(inbound(1, MANAGE_GAMES))

    labels = [button.label for button in games.buttons if button.value.startswith("ag:manage:")]
    assert labels == ["Closed but managed"]
    assert ARCHIVED_GAMES in {button.label for button in games.buttons}
    assert sum(button.label == ARCHIVED_GAMES for button in games.buttons) == 1

    archive = await engine.handle(inbound(2, ARCHIVED_GAMES, callback=True))
    assert [button.label for button in archive.buttons if button.value.startswith("ag:manage:")] == [
        f"Archived {index}" for index in range(PAGE_SIZE, 0, -1)
    ]
    assert PAGE_BACK not in {button.label for button in archive.buttons}
    archive_next = next(button.value for button in archive.buttons if button.label == PAGE_FORWARD)
    assert archive.buttons[-1] == Button(label="Отмена", value=MANAGE_GAMES)

    archive_last = await engine.handle(inbound(3, archive_next, callback=True))
    assert [button.label for button in archive_last.buttons if button.value.startswith("ag:manage:")] == ["Archived 0"]
    assert PAGE_BACK in {button.label for button in archive_last.buttons}
    assert PAGE_FORWARD not in {button.label for button in archive_last.buttons}

    archived_management = await engine.handle(inbound(4, "ag:manage:event-archived-0", callback=True))
    assert "Статус: Архивирована" in archived_management.text
    assert archived_management.buttons == [
        Button(label=UNARCHIVE_GAME, value="ag:unarchive:event-archived-0"),
        Button(label="⬅️ Назад", value=ARCHIVED_GAMES),
    ]

    accepted = await engine.handle(inbound(5, "ag:unarchive:event-archived-0", callback=True))

    assert accepted.text == "⏳ Разархивация принята в обработку."
    assert accepted.deferred
    assert accepted.command_enqueued
    assert accepted.buttons == [Button(label="⬅️ Назад", value=MANAGE_GAMES)]
    assert publisher.commands[-1].operation is Operation.UNARCHIVE_EVENT
    assert publisher.commands[-1].participant_key is None
    assert publisher.commands[-1].reply_context.text_success == (
        "📤 Игра «Archived 0» разархивирована. Регистрация остаётся закрытой."
    )
    assert publisher.commands[-1].reply_context.buttons == [Button(label="⬅️ Назад", value=MANAGE_GAMES)]


@pytest.mark.asyncio
async def test_gamemaster_cannot_list_or_open_archived_games(disk_store: MemoryDiskStore, event: Event) -> None:
    archived = Event.model_validate(
        event.model_dump()
        | {
            "status": EventStatus.CLOSED,
            "archived_at": datetime.now(UTC),
        }
    )
    engine, _, _, _ = await engine_setup(disk_store, archived, gamemaster=True)

    games = await engine.handle(inbound(1, MANAGE_GAMES))
    forged = await engine.handle(inbound(2, f"ag:manage:{event.event_id}", callback=True))
    forged_action = await engine.handle(inbound(3, f"ag:tables:{event.event_id}", callback=True))
    forged_unarchive = await engine.handle(inbound(4, f"ag:unarchive:{event.event_id}", callback=True))

    assert not any(button.value == f"ag:manage:{event.event_id}" for button in games.buttons)
    assert forged.text == "Игра не найдена."
    assert event.name not in forged.text
    assert forged_action.text == "Игра не найдена."
    assert event.name not in forged_action.text
    assert forged_unarchive.text == "Игра не найдена."
    assert event.name not in forged_unarchive.text


@pytest.mark.asyncio
async def test_admin_can_choose_any_of_three_explained_statuses(disk_store: MemoryDiskStore, event: Event) -> None:
    event.status = EventStatus.CREATED
    engine, _, publisher, _ = await engine_setup(disk_store, event, admin=True)

    menu = await engine.handle(inbound(1, ADMIN))
    assert MANAGE_GAMES in [button.value for button in menu.buttons]
    games = await engine.handle(inbound(2, MANAGE_GAMES, callback=True))
    assert any(button.value == f"ag:manage:{event.event_id}" for button in games.buttons)
    management = await engine.handle(inbound(3, f"ag:manage:{event.event_id}", callback=True))
    assert CHANGE_STATUS in [button.label for button in management.buttons]
    status_prompt = await engine.handle(inbound(4, f"ag:status:{event.event_id}", callback=True))
    assert "Текущий Статус: Регистрация" in status_prompt.text
    assert "могут записываться, но ещё не могут подтверждать" in status_prompt.text
    assert "могут записываться и подтверждать" in status_prompt.text
    assert "не могут ни записываться, ни подтверждать" in status_prompt.text
    assert "Закрытую регистрацию можно открыть снова" in status_prompt.text
    assert "игровая сессия уже собрана" in status_prompt.text
    assert [button.label for button in status_prompt.buttons[:3]] == [
        STATUS_REGISTRATION,
        STATUS_CONFIRMATION,
        STATUS_CLOSED,
    ]

    expected_operations = [
        ("admin:status:registration", Operation.OPEN_REGISTRATION),
        ("admin:status:confirmation", Operation.OPEN_CONFIRMATION),
        ("admin:status:closed", Operation.CLOSE_EVENT),
    ]
    update = 5
    for callback, operation in expected_operations:
        if update > 5:
            await engine.handle(inbound(update, f"ag:manage:{event.event_id}", callback=True))
            update += 1
            await engine.handle(inbound(update, f"ag:status:{event.event_id}", callback=True))
            update += 1
        queued = await engine.handle(inbound(update, callback, callback=True))
        update += 1
        if operation is Operation.OPEN_CONFIRMATION:
            assert "DD.MM.YY HH:MM" in queued.text
            queued = await engine.handle(inbound(update, "10.09.26 19:00"))
            update += 1
        assert queued.deferred
        assert publisher.commands[-1].operation is operation
        assert queued.buttons == [Button(label="⬅️ Назад", value=f"ag:manage:{event.event_id}")]
        if operation is Operation.OPEN_CONFIRMATION:
            payload = publisher.commands[-1].payload
            assert isinstance(payload, ConfirmationDeadlinePayload)
            assert payload.deadline.strftime("%d.%m.%y %H:%M") == "10.09.26 19:00"


@pytest.mark.asyncio
async def test_event_master_sees_registration_closing_disclaimer(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, publisher, _ = await engine_setup(disk_store, event, gamemaster=True)
    await engine.events.add_leader(
        event.event_id,
        BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1),
    )

    status_prompt = await engine.handle(inbound(1, f"ag:status:{event.event_id}", callback=True))
    queued = await engine.handle(inbound(2, "admin:status:closed", callback=True))

    assert "Закрытую регистрацию можно открыть снова" in status_prompt.text
    assert "игровая сессия уже собрана" in status_prompt.text
    assert queued.deferred
    assert publisher.commands[-1].operation is Operation.CLOSE_EVENT


@pytest.mark.asyncio
async def test_admin_confirmation_deadline_retries_invalid_input_with_same_options(
    disk_store: MemoryDiskStore, event: Event
) -> None:
    event.status = EventStatus.CREATED
    engine, _, publisher, _ = await engine_setup(disk_store, event, admin=True)

    await engine.handle(inbound(1, f"ag:status:{event.event_id}", callback=True))
    prompt = await engine.handle(inbound(2, "admin:status:confirmation", callback=True))
    invalid = await engine.handle(inbound(3, "9.09.26 19:00"))

    assert "DD.MM.YY HH:MM" in prompt.text
    assert prompt.buttons[0].label == "Ближайший четверг 19:00"
    assert "Некорректная" in invalid.text
    assert invalid.buttons == prompt.buttons
    assert not publisher.commands


@pytest.mark.asyncio
async def test_admin_can_choose_nearest_thursday_deadline(
    monkeypatch: pytest.MonkeyPatch,
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    deadline = datetime(2026, 9, 3, 16, tzinfo=UTC)
    monkeypatch.setattr("larp_bot.application.conversation.closest_thursday_19", lambda: deadline)
    event.status = EventStatus.CREATED
    engine, _, publisher, _ = await engine_setup(disk_store, event, admin=True)

    await engine.handle(inbound(1, f"ag:status:{event.event_id}", callback=True))
    await engine.handle(inbound(2, "admin:status:confirmation", callback=True))
    queued = await engine.handle(inbound(3, "admin:deadline:nearest-thursday", callback=True))

    assert "03.09.26 19:00" in queued.text
    payload = publisher.commands[-1].payload
    assert isinstance(payload, ConfirmationDeadlinePayload)
    assert payload.deadline == deadline


@pytest.mark.asyncio
async def test_admin_has_separate_confirmation_reminder_action(disk_store: MemoryDiskStore, event: Event) -> None:
    event.confirmation_deadline = datetime(2026, 9, 10, 16, tzinfo=UTC)
    engine, _, publisher, _ = await engine_setup(disk_store, event, admin=True)

    menu = await engine.handle(inbound(1, ADMIN))
    assert MANAGE_GAMES in [button.value for button in menu.buttons]
    management = await engine.handle(inbound(2, f"ag:manage:{event.event_id}", callback=True))
    management_labels = [button.label for button in management.buttons]
    assert MANAGE_NOTIFICATIONS in management_labels
    assert SEND_CONFIRMATION_REMINDER not in management_labels
    notifications = await engine.handle(inbound(3, f"ag:notifications:{event.event_id}", callback=True))
    assert SEND_CONFIRMATION_REMINDER in [button.label for button in notifications.buttons]
    queued = await engine.handle(inbound(4, f"ag:remind:{event.event_id}", callback=True))

    assert queued.deferred
    assert publisher.commands[-1].operation is Operation.SEND_CONFIRMATION_REMINDER


@pytest.mark.asyncio
async def test_admin_can_queue_message_for_confirmed_players(disk_store: MemoryDiskStore, event: Event) -> None:
    event.status = EventStatus.CLOSED
    engine, _, publisher, _ = await engine_setup(disk_store, event, admin=True)

    menu = await engine.handle(inbound(1, ADMIN))
    assert MANAGE_GAMES in [button.value for button in menu.buttons]
    management = await engine.handle(inbound(2, f"ag:manage:{event.event_id}", callback=True))
    management_labels = [button.label for button in management.buttons]
    assert MANAGE_NOTIFICATIONS in management_labels
    assert SEND_CONFIRMED_NOTIFICATION not in management_labels
    notifications = await engine.handle(inbound(3, f"ag:notifications:{event.event_id}", callback=True))
    assert [button.label for button in notifications.buttons[:3]] == [
        SEND_CONFIRMATION_REMINDER,
        SEND_CONFIRMED_LINK,
        SEND_CONFIRMED_NOTIFICATION,
    ]
    prompt = await engine.handle(inbound(4, f"ag:notify-message:{event.event_id}", callback=True))
    assert prompt.text.startswith(
        "Эта кнопка нужна для рассылки оргинформации всем подтвердившимся на игру игрокам, "
        "включая подтвердившихся в запас"
    )

    stale_button = await engine.handle(inbound(5, CHANGE_STATUS, callback=True))
    assert "кнопка устарела" in stale_button.text.casefold()
    assert stale_button.buttons == prompt.buttons
    assert not publisher.commands

    preview = await engine.handle(inbound(6, "Сбор в 18:30"))
    assert "Получателей: 0" in preview.text
    assert f"Сообщение от мастеров игры «{event.name}»:\n\nСбор в 18:30" in preview.text
    queued = await engine.handle(inbound(7, "admin:notification:confirm", callback=True))
    assert queued.deferred
    command = publisher.commands[-1]
    assert command.operation is Operation.SEND_CONFIRMED_NOTIFICATION
    assert isinstance(command.payload, NotificationPayload)
    assert command.payload.text == "Сбор в 18:30"


@pytest.mark.asyncio
async def test_admin_link_notification_requires_a_chat_link_and_previews_main_roster_count(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    event.player_amount = 1
    engine, _, publisher, tables = await engine_setup(disk_store, event, admin=True)
    for index in range(2):
        key = f"{index}" * 43
        await tables.enlist(
            event.event_id,
            operation_id=f"enlist-{index}",
            participant_key=key,
            display_name=f"Player {index}",
            wish_play="A",
        )
        await tables.confirm(
            event.event_id,
            operation_id=f"confirm-{index}",
            participant_key=key,
            character_wish="Doctor",
        )

    await engine.handle(inbound(1, f"ag:notifications:{event.event_id}", callback=True))
    prompt = await engine.handle(inbound(2, f"ag:notify-link:{event.event_id}", callback=True))
    assert "только из корректной ссылки на чат Telegram или VK" in prompt.text

    invalid = await engine.handle(inbound(3, "Чат игры: https://vk.me/join/AJQ1d4V/FAtt9xoXEwCAfx1q"))
    assert "Пришлите только корректную ссылку-приглашение" in invalid.text
    assert not publisher.commands

    link = "https://vk.me/join/AJQ1d4V/FAtt9xoXEwCAfx1q"
    preview = await engine.handle(inbound(4, link))
    assert "Получателей: 1" in preview.text
    assert f"Вы в основном составе игры «{event.name}». Вступите в чат: {link}" in preview.text

    queued = await engine.handle(inbound(5, "admin:notification:confirm", callback=True))
    assert queued.deferred
    assert isinstance(publisher.commands[-1].payload, NotificationPayload)
    assert publisher.commands[-1].payload.text == link


@pytest.mark.asyncio
async def test_information_notification_preview_counts_confirmed_reserve_players(
    disk_store: MemoryDiskStore,
    event: Event,
) -> None:
    event.player_amount = 1
    engine, _, _, tables = await engine_setup(disk_store, event, admin=True)
    for index in range(2):
        key = f"{index}" * 43
        await tables.enlist(
            event.event_id,
            operation_id=f"enlist-{index}",
            participant_key=key,
            display_name=f"Player {index}",
            wish_play="A",
        )
        await tables.confirm(
            event.event_id,
            operation_id=f"confirm-{index}",
            participant_key=key,
            character_wish="Doctor",
        )

    await engine.handle(inbound(1, f"ag:notify-message:{event.event_id}", callback=True))
    preview = await engine.handle(inbound(2, "Сбор в 18:30"))

    assert "Получателей: 2" in preview.text


@pytest.mark.asyncio
async def test_event_tables_action_returns_all_three_permanent_links(disk_store: MemoryDiskStore, event: Event) -> None:
    engine, _, _, _ = await engine_setup(disk_store, event, admin=True)

    management = await engine.handle(inbound(1, f"ag:manage:{event.event_id}", callback=True))
    assert EVENT_TABLES in [button.label for button in management.buttons]

    tables = await engine.handle(inbound(2, f"ag:tables:{event.event_id}", callback=True))
    assert tables.text.count("https://disk.example/") == 3
    assert "Публичная таблица (без контактов Telegram и VK)" in tables.text
    assert "Административная таблица" in tables.text
    assert "Таблица пропусков" in tables.text
    assert MASTER_TABLE_DISCLAIMER in tables.text
    assert tables.buttons == [Button(label="⬅️ Назад", value=f"ag:manage:{event.event_id}")]

    repeated = await engine.handle(inbound(3, f"ag:tables:{event.event_id}", callback=True))
    assert repeated.text == tables.text


@pytest.mark.asyncio
@pytest.mark.parametrize("count", [0, 1, 5, 6, 10, 11])
async def test_admin_pagination_uses_shared_page_size(count: int, disk_store: MemoryDiskStore) -> None:
    base = datetime(2026, 1, 1, tzinfo=UTC)
    events = [
        Event(
            event_id=f"event-{index:04d}",
            name=f"Game {index}",
            disk_resource_path=f"disk:/larp-bot/events/event-{index:04d}-game.xlsx",
            public_registration_url=f"https://disk.example/{index}",
            created_at=base + timedelta(seconds=index),
            updated_at=base + timedelta(seconds=index),
        )
        for index in range(count)
    ]
    users = MemoryUserRepository()
    await users.save(
        TelegramUser(
            tg_id=1,
            full_name="Admin",
            vk_url="vk.com/admin",
            crossplay=True,
            larp_experience=True,
            needs_pass=False,
        )
    )
    event_repository = MemoryEventRepository(events)
    tables = MemoryRegistrationRepository()
    showcase = YandexDiskShowcaseRepository(disk_store)
    catalog = RegistrationCatalog(event_repository, tables, showcase)
    engine = ConversationEngine(
        users,
        event_repository,
        RegistrationService(event_repository, catalog, MemoryCommandPublisher(), "secret"),
        EventAdministrationService(event_repository, showcase, catalog, users, "secret"),
        StaticAdminProvider(tg_ids={1}),
    )
    page = await engine.handle(inbound(1, MANAGE_GAMES))
    game_buttons = [button for button in page.buttons if button.value.startswith("ag:manage:")]
    assert len(game_buttons) == min(count, PAGE_SIZE)
    assert [button.label for button in game_buttons] == [
        f"Game {index}" for index in range(count - 1, max(count - PAGE_SIZE - 1, -1), -1)
    ]
    assert "https://disk.example/" not in page.text
    labels = {button.label for button in page.buttons}
    assert (PAGE_FORWARD in labels) is (count > PAGE_SIZE)
    assert PAGE_BACK not in labels
    assert page.buttons[-1] == Button(label="Отмена", value=ADMIN_MENU)


@pytest.mark.asyncio
async def test_player_event_pickers_have_bidirectional_five_item_pages(disk_store: MemoryDiskStore) -> None:
    base = datetime(2026, 1, 1, tzinfo=UTC)
    events = [
        Event(
            event_id=f"event-{index:04d}",
            name=f"Game {index}",
            disk_resource_path=f"disk:/larp-bot/events/event-{index:04d}/game.xlsx",
            public_registration_url=f"https://disk.example/{index}",
            status=EventStatus.CONFIRMATION_OPEN,
            registrations_migrated_at=base,
            created_at=base + timedelta(seconds=index),
            updated_at=base + timedelta(seconds=index),
        )
        for index in range(PAGE_SIZE + 1)
    ]
    engine, _, _, registrations = await engine_setup(disk_store, *events)
    await registrations.import_missing(
        [
            Registration(
                event_id=candidate.event_id,
                participant_key=engine.registrations.key(Platform.TELEGRAM, 1, candidate.event_id),
                display_name="Иван Иванов",
                wish_play="Без пожеланий",
            )
            for candidate in events
        ]
    )

    enlist_first = await engine.handle(inbound(1, ENLIST))
    assert len([button for button in enlist_first.buttons if button.value.startswith("select:enlist:")]) == PAGE_SIZE
    assert PAGE_BACK not in {button.label for button in enlist_first.buttons}
    enlist_second = await engine.handle(
        inbound(
            2,
            next(button.value for button in enlist_first.buttons if button.label == PAGE_FORWARD),
            callback=True,
        )
    )
    assert len([button for button in enlist_second.buttons if button.value.startswith("select:enlist:")]) == 1
    assert PAGE_BACK in {button.label for button in enlist_second.buttons}
    assert PAGE_FORWARD not in {button.label for button in enlist_second.buttons}
    assert enlist_second.buttons[-1] == Button(label=MAIN_MENU_LABEL, value="Отмена")

    registered_first = await engine.handle(inbound(3, CONFIRM))
    registered_second = await engine.handle(
        inbound(
            4,
            next(button.value for button in registered_first.buttons if button.label == PAGE_FORWARD),
            callback=True,
        )
    )
    assert len([button for button in registered_second.buttons if button.value.startswith("select:confirm:")]) == 1
    assert PAGE_BACK in {button.label for button in registered_second.buttons}
    assert PAGE_FORWARD not in {button.label for button in registered_second.buttons}
    returned = await engine.handle(
        inbound(
            5,
            next(button.value for button in registered_second.buttons if button.label == PAGE_BACK),
            callback=True,
        )
    )
    assert len([button for button in returned.buttons if button.value.startswith("select:confirm:")]) == PAGE_SIZE
    assert PAGE_BACK not in {button.label for button in returned.buttons}

    wishes_first = await engine.handle(inbound(6, GAME_WISHES))
    assert len([button for button in wishes_first.buttons if button.value.startswith("select:wishes:")]) == PAGE_SIZE
    wishes_second = await engine.handle(
        inbound(
            7,
            next(button.value for button in wishes_first.buttons if button.label == PAGE_FORWARD),
            callback=True,
        )
    )
    last_game_buttons = [button for button in wishes_second.buttons if button.value.startswith("select:wishes:")]
    assert len(last_game_buttons) == 1
    assert PAGE_BACK in {button.label for button in wishes_second.buttons}
    selected = await engine.handle(inbound(8, last_game_buttons[0].value, callback=True))
    assert selected.text.startswith("Пожелания по игре «")
    assert [button.label for button in selected.buttons] == [CHARACTER_WISHES, CO_PLAYER_WISHES]


def my_games_event(index: int, **changes: object) -> Event:
    created_at = datetime(2026, 1, 1, tzinfo=UTC) + timedelta(seconds=index)
    event_id = f"my-game-{index:04d}"
    return Event(
        event_id=event_id,
        name=f"Game {index}",
        game_date=date(2026, 10, index + 1),
        disk_resource_path=f"disk:/larp-bot/events/{event_id}/master.xlsx",
        public_registration_url=f"https://disk.example/master/{event_id}",
        public_table_resource_path=f"disk:/larp-bot/events/{event_id}/public.xlsx",
        public_table_public_url=f"https://disk.example/public/{event_id}",
        status=EventStatus.CONFIRMATION_OPEN,
        created_at=created_at,
        updated_at=created_at,
    ).model_copy(update=changes)


def my_games_registration(
    engine: ConversationEngine,
    event: Event,
    status: AttendanceStatus = AttendanceStatus.WAITING,
    *,
    user_id: int = 1,
    signup_delay: int = 0,
) -> Registration:
    return Registration(
        event_id=event.event_id,
        participant_key=engine.registrations.key(Platform.TELEGRAM, user_id, event.event_id),
        display_name="Иван Иванов",
        wish_play="Без пожеланий",
        attendance_status=status,
        created_at=event.created_at + timedelta(minutes=signup_delay),
    )


@pytest.mark.asyncio
async def test_my_games_reports_when_there_are_no_games_to_view(disk_store: MemoryDiskStore) -> None:
    cancelled = my_games_event(0)
    archived = my_games_event(1, archived_at=datetime(2026, 2, 1, tzinfo=UTC))
    foreign = my_games_event(2)
    engine, _, _, tables = await engine_setup(disk_store, cancelled, archived, foreign)
    await tables.import_missing(
        [
            my_games_registration(engine, cancelled, AttendanceStatus.CANCELLED),
            my_games_registration(engine, archived, AttendanceStatus.CONFIRMED),
            my_games_registration(engine, foreign, AttendanceStatus.CONFIRMED, user_id=2),
        ]
    )
    await engine.events.add_leader(archived.event_id, BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1))

    response = await engine.handle(inbound(1, MY_GAMES))

    assert response.text == "Сейчас нет игр, доступных для просмотра."
    assert response.buttons == [Button(label=MAIN_MENU_LABEL, value=MAIN_MENU)]


@pytest.mark.asyncio
async def test_single_my_game_opens_its_page_without_a_list(disk_store: MemoryDiskStore) -> None:
    game = my_games_event(0, confirmed_notifications=["https://t.me/+gamechat"])
    engine, _, _, tables = await engine_setup(disk_store, game, my_games_event(1))
    await tables.import_missing([my_games_registration(engine, game)])

    response = await engine.handle(inbound(1, MY_GAMES))

    assert response.text == (
        "Игра: «Game 0»\n"
        "Дата: 01.10.26\n"
        "Ваш статус: Ожидает подтверждения, основной состав\n\n"
        "Таблица участников: https://disk.example/public/my-game-0000"
    )
    assert response.buttons == [Button(label=MAIN_MENU_LABEL, value=MAIN_MENU)]


@pytest.mark.asyncio
async def test_my_games_lists_player_and_master_games_with_status_and_links(disk_store: MemoryDiskStore) -> None:
    chat = "https://t.me/+gamechat"
    confirmed = my_games_event(0, confirmed_notifications=[chat, "Сбор в 18:00", chat])
    led = my_games_event(1, status=EventStatus.CLOSED, confirmed_notifications=[chat])
    reserve = my_games_event(2, player_amount=1, confirmed_notifications=[chat])
    cancelled = my_games_event(3)
    archived = my_games_event(4, archived_at=datetime(2026, 2, 1, tzinfo=UTC))
    undated = my_games_event(5, game_date=None)
    engine, _, _, tables = await engine_setup(disk_store, confirmed, led, reserve, cancelled, archived, undated)
    await tables.import_missing(
        [
            my_games_registration(engine, confirmed, AttendanceStatus.CONFIRMED),
            my_games_registration(engine, reserve, user_id=2),
            my_games_registration(engine, reserve, AttendanceStatus.CONFIRMED, signup_delay=1),
            my_games_registration(engine, cancelled, AttendanceStatus.CANCELLED),
            my_games_registration(engine, archived, AttendanceStatus.CONFIRMED),
            my_games_registration(engine, undated),
        ]
    )
    await engine.events.add_leader(led.event_id, BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1))

    listing = await engine.handle(inbound(1, MY_GAMES))

    assert listing.text == "Выберите игру, чтобы посмотреть информацию о ней:"
    assert listing.buttons == [
        *(Button(label=game.name, value=f"select:mine:{game.event_id}") for game in (confirmed, led, reserve, undated)),
        Button(label=MAIN_MENU_LABEL, value="Отмена"),
    ]

    confirmed_page = await engine.handle(inbound(2, f"select:mine:{confirmed.event_id}", callback=True))
    assert confirmed_page.text == (
        "Игра: «Game 0»\n"
        "Дата: 01.10.26\n"
        "Ваш статус: Подтверждено, основной состав\n\n"
        f"Чат игры: {chat}\n"
        "Таблица участников: https://disk.example/public/my-game-0000"
    )
    assert confirmed_page.buttons == [
        Button(label=BACK, value=MY_GAMES),
        Button(label=MAIN_MENU_LABEL, value=MAIN_MENU),
    ]

    led_page = await engine.handle(inbound(3, f"select:mine:{led.event_id}", callback=True))
    assert led_page.text == (
        "Игра: «Game 1»\n"
        "Дата: 02.10.26\n"
        "Ваш статус: Мастер\n\n"
        "Мастера игры:\nИван Иванов\n\n"
        f"Чат игры: {chat}\n"
        "Таблица участников: https://disk.example/public/my-game-0001"
    )

    reserve_page = await engine.handle(inbound(4, f"select:mine:{reserve.event_id}", callback=True))
    assert reserve_page.text == (
        "Игра: «Game 2»\n"
        "Дата: 03.10.26\n"
        "Ваш статус: Подтверждено, запас\n\n"
        "Таблица участников: https://disk.example/public/my-game-0002"
    )

    undated_page = await engine.handle(inbound(5, f"select:mine:{undated.event_id}", callback=True))
    assert undated_page.text.startswith(
        "Игра: «Game 5»\nДата: не указана\nВаш статус: Ожидает подтверждения, основной состав\n\n"
    )

    returned = await engine.handle(inbound(6, MY_GAMES, callback=True))
    assert returned.buttons == listing.buttons

    no_longer_playing = await engine.handle(inbound(7, f"select:mine:{cancelled.event_id}", callback=True))
    assert no_longer_playing.text == "Вы больше не участвуете в этой игре."
    inactive = await engine.handle(inbound(8, f"select:mine:{archived.event_id}", callback=True))
    assert inactive.text == INACTIVE_GAME_RESPONSE


@pytest.mark.asyncio
async def test_my_games_pages_hold_five_games(disk_store: MemoryDiskStore) -> None:
    games = [my_games_event(index) for index in range(PAGE_SIZE + 1)]
    engine, _, _, tables = await engine_setup(disk_store, *games)
    await tables.import_missing([my_games_registration(engine, game) for game in games[::2]])
    for game in games[1::2]:
        await engine.events.add_leader(game.event_id, BotIdentity(platform=Platform.TELEGRAM, platform_user_id=1))

    first = await engine.handle(inbound(1, MY_GAMES))
    second = await engine.handle(
        inbound(2, next(button.value for button in first.buttons if button.label == PAGE_FORWARD), callback=True)
    )
    returned = await engine.handle(
        inbound(3, next(button.value for button in second.buttons if button.label == PAGE_BACK), callback=True)
    )

    def game_labels(buttons: list[Button]) -> list[str]:
        return [button.label for button in buttons if button.value.startswith("select:mine:")]

    assert game_labels(first.buttons) == [f"Game {index}" for index in range(PAGE_SIZE)]
    assert PAGE_BACK not in {button.label for button in first.buttons}
    assert game_labels(second.buttons) == [f"Game {PAGE_SIZE}"]
    assert PAGE_FORWARD not in {button.label for button in second.buttons}
    assert game_labels(returned.buttons) == game_labels(first.buttons)


@pytest.mark.asyncio
async def test_my_game_page_lists_master_names_and_profiles(disk_store: MemoryDiskStore) -> None:
    game = my_games_event(0)
    engine, users, _, tables = await engine_setup(disk_store, game)
    await tables.import_missing([my_games_registration(engine, game)])
    await users.save(TelegramUser(tg_id=21, full_name="Смирнова Анна Петровна", telegram_handle="@anna_master"))
    await users.save(VkUser(vk_id=42, full_name="Орлов Борис"))
    for platform, master_id in ((Platform.TELEGRAM, 21), (Platform.VK, 42)):
        await engine.events.add_leader(game.event_id, BotIdentity(platform=platform, platform_user_id=master_id))

    response = await engine.handle(inbound(1, MY_GAMES))

    assert response.text == (
        "Игра: «Game 0»\n"
        "Дата: 01.10.26\n"
        "Ваш статус: Ожидает подтверждения, основной состав\n\n"
        "Мастера игры:\n"
        "Орлов Борис — https://vk.com/id42\n"
        "Смирнова Анна — https://t.me/anna_master\n\n"
        "Таблица участников: https://disk.example/public/my-game-0000"
    )
