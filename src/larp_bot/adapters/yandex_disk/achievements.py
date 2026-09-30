from __future__ import annotations

import math
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from threading import Lock
from typing import cast

import pypdfium2 as pdfium
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

PLACEHOLDER_TEXT = "ЛОГОС"
ACHIEVEMENT_TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "assets" / "pdf_achievements_usual.pdf"
RENDER_SCALE = 3
TITLE_BOX_WIDTH_POINTS = 180.0
TITLE_BOX_HEIGHT_POINTS = 24.0
TITLE_HORIZONTAL_PADDING_POINTS = 6.0
MIN_FONT_SIZE_POINTS = 1.0
_FONT_SUBSET_PREFIX = re.compile(r"^[A-Z]{6}\+")
_PDFIUM_LOCK = Lock()


class AchievementTemplateError(RuntimeError):
    pass


@dataclass(frozen=True)
class _Placeholder:
    center_x: float
    center_y: float
    left: float
    top: float
    right: float
    bottom: float
    font_name: str
    font_size: float


def _normalized_game_name(game_name: str) -> str:
    normalized = " ".join(game_name.split())
    if not normalized:
        raise ValueError("achievement game name must not be empty")
    return normalized


def _text_width_em(text: str) -> float:
    """Conservative Oswald-width estimate used to keep the title on one line."""

    narrow = frozenset(".,:;!|ijlіїтгчТГЧІЇ")
    wide = frozenset("mwMWжшщыюФЖШЩЫЮ@%")
    width = 0.0
    for character in text:
        if character.isspace():
            width += 0.32
        elif character in narrow:
            width += 0.38
        elif character in wide:
            width += 0.78
        else:
            width += 0.6
    return max(width, 1.0)


def _fitted_font_size(text: str, original_size: float) -> float:
    available = TITLE_BOX_WIDTH_POINTS - 2 * TITLE_HORIZONTAL_PADDING_POINTS
    fitted = available / _text_width_em(text) * 0.94
    fitted = math.floor(fitted * 4) / 4
    return max(MIN_FONT_SIZE_POINTS, min(original_size, fitted))


def _font_name(raw_name: str) -> str:
    name = _FONT_SUBSET_PREFIX.sub("", raw_name).replace("-", " ").strip()
    return name or "Oswald Light"


def _placeholder(text_page: pdfium.PdfTextPage, page_height: float, index: int, count: int) -> _Placeholder:
    rectangle_count = text_page.count_rects(index, count)
    if rectangle_count < 1:
        raise AchievementTemplateError("achievement placeholder has no visible bounds")
    rectangles = [text_page.get_rect(rectangle_index) for rectangle_index in range(rectangle_count)]
    left = min(rectangle[0] for rectangle in rectangles)
    pdf_bottom = min(rectangle[1] for rectangle in rectangles)
    right = max(rectangle[2] for rectangle in rectangles)
    pdf_top = max(rectangle[3] for rectangle in rectangles)
    text_object = text_page.get_textobj(index)
    if text_object is None:
        raise AchievementTemplateError("achievement placeholder is not a PDF text object")
    matrix = text_object.get_matrix()
    vertical_scale = math.hypot(matrix.c, matrix.d)
    font_size = text_object.get_font_size() * vertical_scale
    font_name = _font_name(text_object.get_font().get_base_name())
    top = page_height - pdf_top
    bottom = page_height - pdf_bottom
    return _Placeholder(
        center_x=(left + right) / 2,
        center_y=(top + bottom) / 2,
        left=left,
        top=top,
        right=right,
        bottom=bottom,
        font_name=font_name,
        font_size=font_size,
    )


def _render_background(page: pdfium.PdfPage, placeholders: list[_Placeholder]) -> Image.Image:
    bitmap = page.render(scale=RENDER_SCALE, optimize_mode="print")
    try:
        image = bitmap.to_pil().convert("RGB").copy()
    finally:
        bitmap.close()
    drawing = ImageDraw.Draw(image)
    # PDFium's search rectangle is tighter than the glyph ascenders and
    # descenders, so erase a slightly larger area without touching the card.
    padding = 8 * RENDER_SCALE
    for placeholder in placeholders:
        drawing.rectangle(
            (
                math.floor(placeholder.left * RENDER_SCALE) - padding,
                math.floor(placeholder.top * RENDER_SCALE) - padding,
                math.ceil(placeholder.right * RENDER_SCALE) + padding,
                math.ceil(placeholder.bottom * RENDER_SCALE) + padding,
            ),
            fill="white",
        )
    return cast(Image.Image, image)


def _presentation_bytes(
    page_width: float,
    page_height: float,
    background: Image.Image,
    placeholders: list[_Placeholder],
    game_name: str,
) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Pt(page_width)
    presentation.slide_height = Pt(page_height)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    background_bytes = BytesIO()
    background.save(background_bytes, format="PNG", compress_level=6)
    background_bytes.seek(0)
    slide.shapes.add_picture(background_bytes, 0, 0, width=Pt(page_width), height=Pt(page_height))

    for placeholder in placeholders:
        text_box = slide.shapes.add_textbox(
            Pt(placeholder.center_x - TITLE_BOX_WIDTH_POINTS / 2),
            Pt(placeholder.center_y - TITLE_BOX_HEIGHT_POINTS / 2),
            Pt(TITLE_BOX_WIDTH_POINTS),
            Pt(TITLE_BOX_HEIGHT_POINTS),
        )
        frame = text_box.text_frame
        frame.clear()
        frame.word_wrap = False
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = Pt(TITLE_HORIZONTAL_PADDING_POINTS)
        frame.margin_right = Pt(TITLE_HORIZONTAL_PADDING_POINTS)
        frame.margin_top = 0
        frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = game_name
        run.font.name = placeholder.font_name
        run.font.size = Pt(_fitted_font_size(game_name, placeholder.font_size))
        run.font.bold = False
        run.font.color.rgb = RGBColor(0, 0, 0)  # type: ignore[no-untyped-call]

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _render_presentation(template: bytes, game_name: str) -> bytes:
    try:
        with pdfium.PdfDocument(template) as document:
            if len(document) != 1:
                raise AchievementTemplateError("achievement template must contain exactly one page")
            page = document[0]
            try:
                page_width, page_height = page.get_size()
                text_page = page.get_textpage()
                try:
                    search = text_page.search(PLACEHOLDER_TEXT, match_case=True, match_whole_word=True)
                    try:
                        placeholders: list[_Placeholder] = []
                        while match := search.get_next():
                            placeholders.append(_placeholder(text_page, page_height, *match))
                    finally:
                        search.close()
                finally:
                    text_page.close()
                if not placeholders:
                    raise AchievementTemplateError("achievement template has no ЛОГОС placeholders")
                background = _render_background(page, placeholders)
            finally:
                page.close()
    except pdfium.PdfiumError as error:
        raise AchievementTemplateError("achievement template is not a valid PDF") from error
    return _presentation_bytes(page_width, page_height, background, placeholders, game_name)


def achievement_presentation_bytes(template: bytes, game_name: str) -> bytes:
    """Render the PDF template as a PPTX while leaving the source bytes untouched."""

    normalized_name = _normalized_game_name(game_name)
    # PDFium is process-global and not thread-safe. Achievement rendering runs
    # in worker threads, so serialize PDF operations within a warm function.
    with _PDFIUM_LOCK:
        return _render_presentation(template, normalized_name)
