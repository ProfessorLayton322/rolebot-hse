"""Edit only the current season's OOXML, retaining historical sheets and caches.

The supplied Google Sheets export contains cached DUMMYFUNCTION/array formulas.
An openpyxl round trip discards those caches in *every* season. Keeping the ZIP
parts intact also preserves styles, checkbox metadata and unrelated worksheets.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from copy import deepcopy
from io import BytesIO
from typing import ClassVar
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from openpyxl.utils import column_index_from_string, get_column_letter

from larp_bot.application.services import DomainError
from larp_bot.domain.models import parse_full_name

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
ET.register_namespace("", NS)
ET.register_namespace("r", REL)


def tag(name: str) -> str:
    return f"{{{NS}}}{name}"


def normalized_name(value: str) -> str:
    return " ".join(value.split()).casefold().replace("ё", "е")


def reversed_name(value: str) -> str | None:
    try:
        parts = parse_full_name(value, alphabet="cyrillic")
    except ValueError:
        return None
    if len(parts) != 2:
        return None
    return normalized_name(f"{parts[1]} {parts[0]}")


class StatisticsWorkbook:
    required_headers: ClassVar[dict[int, str]] = {1: "Игрок", 4: "5 игр", 13: "50 игр"}
    fixed_column = 13
    game_column = 14
    carried_column = 3
    total_column = 2
    award_column = 4
    footer_titles: ClassVar[tuple[str, ...]] = (
        "Всего",
        "Участники сезона",
        "Новички",
        "Активные игроки сезона",
    )
    layout_error = "Нужна XLSX-таблица с сезонным листом, колонками Игрок, итоги и 5–50 игр."
    skip_nonblank_game_formatting = True
    match_reversed_names = True

    def __init__(self, content: bytes) -> None:
        self.original = content
        try:
            with ZipFile(BytesIO(content)) as archive:
                self.parts = {name: archive.read(name) for name in archive.namelist()}
            self.book = ET.fromstring(self.parts["xl/workbook.xml"])  # noqa: S314
            self.rels = ET.fromstring(self.parts["xl/_rels/workbook.xml.rels"])  # noqa: S314
            sheets = self.book.find(tag("sheets"))
            assert sheets is not None
            self.sheets = sheets
            seasons = [s for s in self.sheets if re.fullmatch(r"20\d\d-20\d\d", s.attrib["name"])]
            self.season = max(seasons, key=lambda s: s.attrib["name"])
            target = next(r.attrib["Target"] for r in self.rels if r.attrib["Id"] == self.season.attrib[f"{{{REL}}}id"])
            self.path = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
            self.sheet = ET.fromstring(self.parts[self.path])  # noqa: S314
            data = self.sheet.find(tag("sheetData"))
            assert data is not None
            self.data = data
            self.strings: list[str] = []
            if "xl/sharedStrings.xml" in self.parts:
                strings = ET.fromstring(self.parts["xl/sharedStrings.xml"])  # noqa: S314
                self.strings = ["".join(t.text or "" for t in s.iter(tag("t"))) for s in strings]
            self.rows = {int(r.attrib["r"]): r for r in self.data}
            self.cells = {c.attrib["r"]: c for r in self.data for c in r}
            if any(self.value(f"{get_column_letter(col)}1") != value for col, value in self.required_headers.items()):
                raise ValueError("unsupported layout")
            self.last_column = max(
                column_index_from_string(re.sub(r"\d", "", c.attrib["r"]))
                for c in self.rows[1]
                if self.value(c.attrib["r"]) not in (None, "")
            )
            self.last_player = 1
            for row in sorted(self.rows):
                if row == 1:
                    continue
                name = self.value(f"A{row}")
                cell = self.cells.get(f"A{row}")
                if not isinstance(name, str) or name in self.footer_titles:
                    break
                if cell is not None and cell.find(tag("f")) is not None:
                    break
                self.last_player = row
            footer_labels = set(self.footer_titles)
            for index, xml_row in self.rows.items():
                if index <= self.last_player:
                    continue
                populated = any(
                    self.value(c.attrib["r"]) not in (None, "") or c.find(tag("f")) is not None for c in xml_row
                )
                if not populated:
                    continue
                name = self.value(f"A{index}")
                cell = self.cells.get(f"A{index}")
                exported_total = index == self.last_player + 1 and cell is not None and cell.find(tag("f")) is not None
                if index > self.last_player + len(self.footer_titles) or (
                    name not in footer_labels and not exported_total
                ):
                    raise ValueError("unexpected data below player rows")
            style_game_column = self.game_column if self.last_column >= self.game_column else self.fixed_column
            self.player_styles = {col: self.style(2, col) for col in range(1, self.fixed_column + 1)}
            self.game_header_style = self.style(1, style_game_column)
            self.game_cell_style = self.style(2, style_game_column)
            summary = self.last_player + 1
            self.footer_styles = {
                (offset, col): self.style(summary + offset, col)
                for offset in range(len(self.footer_titles))
                for col in range(1, self.last_column + 1)
                if f"{get_column_letter(col)}{summary + offset}" in self.cells
            }
        except (BadZipFile, KeyError, ValueError, AssertionError, StopIteration, ET.ParseError) as exc:
            raise DomainError(self.layout_error) from exc

    def value(self, address: str) -> str | float | bool | None:
        cell = self.cells.get(address)
        if cell is None:
            return None
        kind = cell.get("t")
        value = cell.findtext(tag("v"))
        if kind == "inlineStr":
            return "".join(t.text or "" for t in cell.iter(tag("t")))
        if value is None:
            return None
        if kind == "s":
            return self.strings[int(value)]
        if kind in {"str", "e"}:
            return value
        if kind == "b":
            return value == "1"
        return float(value)

    def number(self, address: str, *, sum_cell: bool = False) -> float:
        value = self.value(address)
        cell = self.cells.get(address)
        if isinstance(value, (float, bool)):
            return float(value)
        # SUM(range) ignores text markers such as the supplied sheet's ❌.
        if sum_cell and isinstance(value, str) and (cell is None or cell.get("t") != "e"):
            return 0
        if value in (None, "") and (cell is None or cell.find(tag("f")) is None):
            return 0
        raise DomainError(f"Нет числового результата в {self.season.attrib['name']}!{address}. Пересчитайте XLSX.")

    def total(self, row: int) -> float:
        return self.number(f"{get_column_letter(self.carried_column)}{row}") + sum(
            self.number(f"{get_column_letter(c)}{row}", sum_cell=True)
            for c in range(self.game_column, self.last_column + 1)
        )

    def set_cell(
        self, row: int, col: int, value: str | float | bool, *, style: str | None = None, formula: str | None = None
    ) -> None:
        address = f"{get_column_letter(col)}{row}"
        if row not in self.rows:
            self.rows[row] = ET.SubElement(self.data, tag("row"), r=str(row))
        cell = self.cells.get(address)
        if cell is None:
            cell = ET.SubElement(self.rows[row], tag("c"), r=address)
            self.cells[address] = cell
        if style is not None:
            cell.set("s", style)
        for child in list(cell):
            cell.remove(child)
        cell.attrib.pop("t", None)
        if formula:
            ET.SubElement(cell, tag("f")).text = formula
        if isinstance(value, str):
            cell.set("t", "inlineStr")
            ET.SubElement(ET.SubElement(cell, tag("is")), tag("t")).text = value
        else:
            if isinstance(value, bool):
                cell.set("t", "b")
            ET.SubElement(cell, tag("v")).text = str(int(value) if isinstance(value, bool) else value)

    def set_blank_cell(self, row: int, col: int, *, style: str) -> None:
        address = f"{get_column_letter(col)}{row}"
        if row not in self.rows:
            self.rows[row] = ET.SubElement(self.data, tag("row"), r=str(row))
        cell = self.cells.get(address)
        if cell is None:
            cell = ET.SubElement(self.rows[row], tag("c"), r=address)
            self.cells[address] = cell
        cell.set("s", style)
        cell.attrib.pop("t", None)
        for child in list(cell):
            cell.remove(child)

    def style(self, row: int, col: int) -> str:
        cell = self.cells.get(f"{get_column_letter(col)}{row}")
        return cell.get("s", "0") if cell is not None else "0"

    def has_game(self, name: str) -> bool:
        return any(
            normalized_name(str(self.value(f"{get_column_letter(c)}1") or "")) == normalized_name(name)
            for c in range(self.game_column, self.last_column + 1)
        )

    def has_any_player(self, names: Sequence[str]) -> bool:
        candidates = {normalized_name(name) for name in names}
        for row in range(2, self.last_player + 1):
            table_name = str(self.value(f"A{row}") or "")
            if normalized_name(table_name) in candidates:
                return True
            if self.match_reversed_names and reversed_name(table_name) in candidates:
                return True
        return False

    @property
    def season_name(self) -> str:
        return self.season.attrib["name"]

    def add_game(self, season_year: int, name: str, players: Sequence[str]) -> bytes | None:
        """Add a game to the requested season, creating that season when needed."""
        title = f"{season_year}-{season_year + 1}"
        if self.season_name != title:
            created = self.new_season(season_year)
            if created is None:
                raise DomainError("Текущий сезон не является последним листом таблицы.")
        if self.has_game(name):
            return None
        return self.mark_game(name, players)

    def carried_header(self, year: int) -> str:
        return f"Игр на 01.09.{year}"

    def new_player_value(self, col: int, name: str) -> str | float | bool:
        if col == 1:
            return name
        if col >= self.award_column:
            return False
        return 0

    def footer_style(self, offset: int, col: int) -> str:
        exact = self.footer_styles.get((offset, col))
        if exact is not None:
            return exact
        if col >= self.game_column:
            game_style = self.footer_styles.get((offset, self.game_column))
            if game_style is not None:
                return game_style
        return self.style(1, min(col, self.fixed_column))

    def prepare_footer(self, summary: int) -> None:
        for offset in range(len(self.footer_titles)):
            for col in range(1, self.last_column + 1):
                self.set_blank_cell(summary + offset, col, style=self.footer_style(offset, col))

    def new_season(self, year: int) -> bytes | None:
        title = f"{year}-{year + 1}"
        if any(s.attrib["name"] == title for s in self.sheets):
            return None
        if title < self.season.attrib["name"]:
            raise DomainError("В таблице уже есть более поздний сезон.")
        totals = {r: self.total(r) for r in range(2, self.last_player + 1)}
        # Start with the fixed columns only, carrying names, totals and award flags.
        for xml_row in list(self.data):
            index = int(xml_row.attrib["r"])
            if index > self.last_player:
                self.data.remove(xml_row)
                self.rows.pop(index, None)
            else:
                for cell in list(xml_row):
                    if column_index_from_string(re.sub(r"\d", "", cell.attrib["r"])) > self.fixed_column:
                        xml_row.remove(cell)
        self.cells = {c.attrib["r"]: c for r in self.data for c in r}
        self.set_cell(1, self.carried_column, self.carried_header(year))
        for row, total in totals.items():
            self.set_cell(row, self.carried_column, total)
        self.last_column = self.fixed_column
        # A new sheet must not inherit relationships to old comments/drawings.
        for node in list(self.sheet):
            if node.tag not in {
                tag(n)
                for n in (
                    "sheetPr",
                    "dimension",
                    "sheetViews",
                    "sheetFormatPr",
                    "cols",
                    "sheetData",
                    "conditionalFormatting",
                    "pageMargins",
                    "pageSetup",
                    "printOptions",
                )
            }:
                self.sheet.remove(node)
        columns = self.sheet.find(tag("cols"))
        if columns is not None:
            for col in list(columns):
                if int(col.attrib["min"]) > self.fixed_column:
                    columns.remove(col)
                elif int(col.attrib["max"]) > self.fixed_column:
                    col.set("max", str(self.fixed_column))
        index = max(int(s.attrib["sheetId"]) for s in self.sheets) + 1
        self.path = f"xl/worksheets/statsSeason{index}.xml"
        rid = f"statsSeason{index}"
        ET.SubElement(
            self.rels,
            f"{{{PKG}}}Relationship",
            Id=rid,
            Type=f"{REL}/worksheet",
            Target=f"worksheets/statsSeason{index}.xml",
        )
        self.sheets.insert(0, ET.Element(tag("sheet"), name=title, sheetId=str(index), attrib={f"{{{REL}}}id": rid}))
        types = ET.fromstring(self.parts["[Content_Types].xml"])  # noqa: S314
        ET.SubElement(
            types,
            f"{{{CT}}}Override",
            PartName=f"/{self.path}",
            ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
        )
        self.parts["[Content_Types].xml"] = ET.tostring(types)
        self.parts["xl/_rels/workbook.xml.rels"] = ET.tostring(self.rels)
        for view in self.book.iter(tag("workbookView")):
            view.set("activeTab", "0")
        self.parts["xl/workbook.xml"] = ET.tostring(self.book)
        self.refresh()
        return self.save()

    def mark_game(self, name: str, players: Sequence[str]) -> bytes | None:
        if self.has_game(name):
            return None
        lookup: dict[str, list[int]] = {}
        reversed_lookup: dict[str, list[int]] = {}
        for row in range(2, self.last_player + 1):
            table_name = str(self.value(f"A{row}"))
            lookup.setdefault(normalized_name(table_name), []).append(row)
            reversed_key = reversed_name(table_name) if self.match_reversed_names else None
            if reversed_key is not None:
                reversed_lookup.setdefault(reversed_key, []).append(row)
        self.last_column += 1
        self.set_cell(1, self.last_column, name, style=self.game_header_style)
        # Rebuild the known summary rows below the extended participant list.
        for xml_row in list(self.data):
            index = int(xml_row.attrib["r"])
            if index > self.last_player:
                self.data.remove(xml_row)
                self.rows.pop(index, None)
                for cell in xml_row:
                    self.cells.pop(cell.attrib["r"], None)
        for player in players:
            key = normalized_name(player)
            matching_rows = lookup.get(key)
            if matching_rows is None:
                matching_rows = reversed_lookup.get(key)
            if matching_rows is None:
                self.last_player += 1
                row = self.last_player
                matching_rows = lookup[key] = [row]
                for col in range(1, self.fixed_column + 1):
                    self.set_cell(row, col, self.new_player_value(col, player), style=self.player_styles[col])
            for row in matching_rows:
                self.set_cell(row, self.last_column, 1, style=self.game_cell_style)
        self.refresh()
        return self.save()

    def refresh(self) -> None:
        last = self.last_player
        end = get_column_letter(self.last_column)
        totals = {row: self.total(row) for row in range(2, last + 1)}
        for row, total in totals.items():
            formula = (
                f"C{row}+SUM({get_column_letter(self.game_column)}{row}:{end}{row})"
                if self.last_column >= self.game_column
                else f"C{row}"
            )
            self.set_cell(row, self.total_column, total, formula=formula)
        summary = last + 1
        self.prepare_footer(summary)
        for offset, title in enumerate(self.footer_titles):
            self.set_cell(summary + offset, 1, title, style=self.footer_style(offset, 1))
        for col in (2, 3):
            letter = get_column_letter(col)
            count = sum(self.number(f"{letter}{row}") > 0 for row in totals)
            self.set_cell(
                summary,
                col,
                count,
                style=self.footer_style(0, col),
                formula=f'COUNTIF({letter}2:{letter}{last},">0")',
            )
        for col in range(self.game_column, self.last_column + 1):
            letter = get_column_letter(col)
            attendance_count = sum(self.number(f"{letter}{row}", sum_cell=True) for row in totals)
            self.set_cell(
                summary,
                col,
                attendance_count,
                style=self.footer_style(0, col),
                formula=f"SUM({letter}2:{letter}{last})",
            )
        for offset, threshold in ((1, 0), (3, 3)):
            count = sum(total > self.number(f"C{row}") + threshold for row, total in totals.items())
            self.set_cell(
                summary + offset,
                2,
                count,
                style=self.footer_style(offset, 2),
                formula=f"SUMPRODUCT(--(B2:B{last}>C2:C{last}+{threshold}))",
            )
        self.set_cell(
            summary + 2,
            2,
            self.number(f"B{summary}") - self.number(f"C{summary}"),
            style=self.footer_style(2, 2),
            formula=f"B{summary}-C{summary}",
        )
        self.refresh_conditional_formatting(last, end)
        dimension = self.sheet.find(tag("dimension"))
        if dimension is not None:
            dimension.set("ref", f"A1:{end}{summary + len(self.footer_titles) - 1}")

    def refresh_conditional_formatting(self, last: int, end: str) -> None:
        """Retain rule expressions and differential styles while extending their ranges."""

        game_letter = get_column_letter(self.game_column)
        total_letter = get_column_letter(self.total_column)
        for cf in list(self.sheet.findall(tag("conditionalFormatting"))):
            old = cf.get("sqref", "")
            first = old.split()[0].split(":", 1)[0] if old else ""
            match = re.fullmatch(r"([A-Z]+)(\d+)", first)
            if match is None:
                continue
            column, _row = match.groups()
            if (
                column == game_letter
                and self.skip_nonblank_game_formatting
                and any(rule.get("type") == "notContainsBlanks" for rule in cf)
            ):
                # The template limits this extra green rule to its first two
                # games. Extending it would also paint text such as ❌ green.
                continue
            if column == "A":
                cf.set("sqref", f"A2:A{last}")
            elif column == total_letter:
                cf.set("sqref", f"{total_letter}1:{total_letter}{last}")
            elif column != game_letter:
                continue
            elif self.last_column < self.game_column:
                # Keep dormant formatting for the first future game without adding cells.
                cf.set("sqref", f"{game_letter}2:{game_letter}{last}")
                for cf_formula in cf.iter(tag("formula")):
                    cf_formula.text = re.sub(r"\b[A-Z]+2\b", f"{game_letter}2", cf_formula.text or "")
            else:
                cf.set("sqref", f"{game_letter}2:{end}{last}")
                for cf_formula in cf.iter(tag("formula")):
                    cf_formula.text = re.sub(r"\b[A-Z]+2\b", f"{game_letter}2", cf_formula.text or "")
        if self.last_column >= self.game_column and not any(
            cf.get("sqref", "").startswith(f"{game_letter}2:")
            for cf in self.sheet.findall(tag("conditionalFormatting"))
        ):
            # New seasons obtain the same green/yellow attendance rules as their source.
            source = ET.fromstring(self.parts[self._source_path()])  # noqa: S314
            for cf in source.findall(tag("conditionalFormatting")):
                if cf.get("sqref", "").startswith(game_letter):
                    clone = deepcopy(cf)
                    clone.set("sqref", f"{game_letter}2:{end}{last}")
                    self._insert_cf(clone)

    def _source_path(self) -> str:
        # Current sheet for a game update; original latest season for a new one.
        target = next(r.attrib["Target"] for r in self.rels if r.attrib["Id"] == self.season.attrib[f"{{{REL}}}id"])
        return target.lstrip("/") if target.startswith("/") else f"xl/{target}"

    def _insert_cf(self, rule: ET.Element) -> None:
        children = list(self.sheet)
        following = [i for i, c in enumerate(children) if c.tag in {tag("pageMargins"), tag("pageSetup")}]
        self.sheet.insert(min(following) if following else len(children), rule)

    def save(self) -> bytes:
        self.data[:] = sorted(self.data, key=lambda row: int(row.attrib["r"]))
        for row in self.data:
            row[:] = sorted(row, key=lambda c: column_index_from_string(re.sub(r"\d", "", c.attrib["r"])))
        self.parts[self.path] = ET.tostring(self.sheet, encoding="utf-8", xml_declaration=True)
        output = BytesIO()
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
            for name, content in self.parts.items():
                archive.writestr(name, content)
        return output.getvalue()


class MasterStatisticsWorkbook(StatisticsWorkbook):
    """Edit the separate gamemaster workbook without rewriting historical sheets."""

    required_headers: ClassVar[dict[int, str]] = {
        1: "Мастер",
        2: "Участие в сезоне",
        3: "Игр на начало сезона",
        4: "Игр на текущий момент",
        5: "1 игра",
        6: "5 игр",
        12: "35 игр",
    }
    fixed_column = 12
    game_column = 13
    carried_column = 3
    total_column = 4
    award_column = 5
    footer_titles: ClassVar[tuple[str, ...]] = ("Всего", "Участники сезона", "Новички")
    layout_error = "Нужна XLSX-таблица мастеров с сезонным листом, итогами, отметками 1–35 игр и колонками игр."
    skip_nonblank_game_formatting = False
    match_reversed_names = False

    def carried_header(self, year: int) -> str:
        return "Игр на начало сезона"

    def refresh(self) -> None:
        last = self.last_player
        end = get_column_letter(self.last_column)
        game_letter = get_column_letter(self.game_column)
        totals = {row: self.total(row) for row in range(2, last + 1)}
        season_totals: dict[int, float] = {}
        for row, total in totals.items():
            carried = self.number(f"C{row}")
            season_total = total - carried
            season_totals[row] = season_total
            total_formula = (
                f"C{row}+SUM({game_letter}{row}:{end}{row})" if self.last_column >= self.game_column else f"C{row}"
            )
            self.set_cell(row, 2, season_total, formula=f"D{row}-C{row}")
            self.set_cell(row, 4, total, formula=total_formula)

        summary = last + 1
        self.prepare_footer(summary)
        for offset, title in enumerate(self.footer_titles):
            self.set_cell(summary + offset, 1, title, style=self.footer_style(offset, 1))
        self.set_cell(
            summary,
            2,
            len(totals),
            style=self.footer_style(0, 2),
            formula=f'COUNTIF(B2:B{last},">=0")',
        )
        for col in range(self.game_column, self.last_column + 1):
            letter = get_column_letter(col)
            attendance_count = sum(self.number(f"{letter}{row}", sum_cell=True) for row in totals)
            self.set_cell(
                summary,
                col,
                attendance_count,
                style=self.footer_style(0, col),
                formula=f"SUM({letter}2:{letter}{last})",
            )
        self.set_cell(
            summary + 1,
            2,
            sum(total > 0 for total in season_totals.values()),
            style=self.footer_style(1, 2),
            formula=f'COUNTIF(B2:B{last},">0")',
        )
        self.set_cell(
            summary + 2,
            2,
            sum(self.number(f"C{row}") == 0 for row in totals),
            style=self.footer_style(2, 2),
            formula=f'COUNTIF(C2:C{last},"=0")',
        )
        self.refresh_conditional_formatting(last, end)
        dimension = self.sheet.find(tag("dimension"))
        if dimension is not None:
            dimension.set("ref", f"A1:{end}{summary + len(self.footer_titles) - 1}")
